"""Lê o arquivo exportado pelo equipamento em qualquer formato que o gerencial aceita.

O formato é reconhecido pelo CONTEÚDO, não pela extensão (o JBV pode gravar ".xls" que por dentro é
outra coisa):
  * .xlsx (zip)                       -> openpyxl
  * .xls antigo (Excel 97-2003, OLE2) -> xlrd
  * HTML ou XML do Excel com extensão .xls, e CSV/TSV -> lidos como texto
Devolve as abas como (nome, linhas), cada linha uma tupla de valores, como o openpyxl com values_only=True.
Número gravado como texto ("12,5", "1.234,5") vira número, com a mesma regra do gerencial
(numero() em src/importador/planilhaAmostras.ts): os dois lados precisam ler as mesmas leituras.
"""
from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
from html.parser import HTMLParser

from openpyxl import load_workbook

Linhas = list[tuple]
Abas = list[tuple[str, Linhas]]

_NUMERO = re.compile(r"^[+-]?(\d+([.,]\d*)?|[.,]\d+)([eE][+-]?\d+)?$")
_MILHAR = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$")


def _numero(v):
    """Texto que é número -> float ("1.234,5" -> 1234.5; "12,5" -> 12.5); vazio -> None; o resto fica como veio."""
    if not isinstance(v, str):
        return v
    s = re.sub(r"\s", "", v)
    if not s:
        return None
    if _MILHAR.match(s):
        return float(s.replace(".", "").replace(",", "."))
    if _NUMERO.match(s):
        return float(s.replace(",", "."))
    return v


def _normalizar(abas: Abas) -> Abas:
    return [(nome, [tuple(_numero(v) for v in l) for l in linhas]) for nome, linhas in abas]


def _xlsx(conteudo: bytes) -> Abas:
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    return [(ws.title, list(ws.iter_rows(values_only=True))) for ws in wb.worksheets]


def _xls(conteudo: bytes) -> Abas:
    import xlrd  # só quando precisa

    livro = xlrd.open_workbook(file_contents=conteudo)
    abas: Abas = []
    for ws in livro.sheets():
        linhas = []
        for r in range(ws.nrows):
            linha = []
            for c in ws.row(r):
                if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    linha.append(None)
                else:
                    linha.append(c.value)
            linhas.append(tuple(linha))
        abas.append((ws.name, linhas))
    return abas


class _Tabelas(HTMLParser):
    """Tabelas de um HTML (Excel/JBV "salvar como .xls" em HTML)."""

    def __init__(self):
        super().__init__()
        self.abas: list[Linhas] = []
        self._linha: list | None = None
        self._celula: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.abas.append([])
        elif tag == "tr":
            self._linha = []
        elif tag in ("td", "th"):
            self._celula = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._celula is not None and self._linha is not None:
            self._linha.append("".join(self._celula).strip() or None)
            self._celula = None
        elif tag == "tr" and self._linha is not None:
            if not self.abas:
                self.abas.append([])
            self.abas[-1].append(tuple(self._linha))
            self._linha = None

    def handle_data(self, data):
        if self._celula is not None:
            self._celula.append(data)


def _html(texto: str) -> Abas:
    p = _Tabelas()
    p.feed(texto)
    return [(f"Tabela{i + 1}", a) for i, a in enumerate(p.abas) if a]


def _xml_excel(texto: str) -> Abas:
    """XML Spreadsheet 2003 (urn:schemas-microsoft-com:office:spreadsheet)."""
    ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
    raiz = ET.fromstring(texto.encode("utf-8"))
    abas: Abas = []
    for i, ws in enumerate(raiz.findall("ss:Worksheet", ns)):
        linhas = []
        for row in ws.findall("ss:Table/ss:Row", ns):
            linha: list = []
            for cel in row.findall("ss:Cell", ns):
                idx = cel.get(f"{{{ns['ss']}}}Index")
                if idx:  # células vazias puladas
                    linha.extend([None] * (int(idx) - 1 - len(linha)))
                d = cel.find("ss:Data", ns)
                linha.append((d.text or None) if d is not None else None)
            linhas.append(tuple(linha))
        abas.append((ws.get(f"{{{ns['ss']}}}Name") or f"Planilha{i + 1}", linhas))
    return abas


def _csv(texto: str) -> Abas:
    amostra = texto[:20000]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";\t,")
    except csv.Error:
        dialeto = csv.excel_tab if amostra.count("\t") > amostra.count(";") else csv.excel
        if amostra.count(";") > amostra.count("\t"):
            dialeto.delimiter = ";"  # type: ignore[misc]
    return [("csv", [tuple(l) for l in csv.reader(io.StringIO(texto), dialeto)])]


def _texto(conteudo: bytes) -> str:
    if conteudo.startswith((b"\xff\xfe", b"\xfe\xff")):
        return conteudo.decode("utf-16")
    for cod in ("utf-8-sig", "cp1252"):
        try:
            return conteudo.decode(cod)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1")


def ler_abas(conteudo: bytes) -> Abas:
    """Todas as abas do arquivo, cada uma como (nome, linhas)."""
    if conteudo[:2] == b"PK":
        return _normalizar(_xlsx(conteudo))
    if conteudo[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _normalizar(_xls(conteudo))
    texto = _texto(conteudo).lstrip()
    if not texto:
        raise ValueError("o arquivo está vazio")
    if texto.startswith("<"):
        if "urn:schemas-microsoft-com:office:spreadsheet" in texto[:5000] and "<html" not in texto[:2000].lower():
            return _normalizar(_xml_excel(texto))
        abas = _html(texto)
        if not abas:
            raise ValueError("o arquivo parece HTML, mas não tem nenhuma tabela")
        return _normalizar(abas)
    return _normalizar(_csv(texto))
