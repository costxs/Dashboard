import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

AQUI = Path(__file__).parent
PEDIDO = json.loads((AQUI / "exemplo_pedido.json").read_text())


@pytest.fixture
def cliente(monkeypatch):
    monkeypatch.setenv("MODO_DEV", "1")
    from app import app
    return TestClient(app)


def pedido(**config):
    p = json.loads(json.dumps(PEDIDO))
    p["config"].update(config)
    return p


@pytest.mark.parametrize("modo", ["por_vazao", "painel", "combinado", "darcy"])
@pytest.mark.parametrize("formato,inicio", [("png", b"\x89PNG"), ("svg", b"<?xml"), ("pdf", b"%PDF")])
def test_todos_os_modos_e_formatos(cliente, modo, formato, inicio):
    r = cliente.post("/permeabilidade/grafico", json=pedido(modo=modo, formato=formato, etapa=2, dpi=100))
    assert r.status_code == 200, r.text
    assert r.content.startswith(inicio)


def test_cor_escolhida_aparece_no_svg(cliente):
    p = pedido(modo="por_vazao", formato="svg", etapa=1)
    p["etapas"][0]["cor"] = "#C8102E"
    r = cliente.post("/permeabilidade/grafico", json=p)
    assert "#c8102e" in r.text.lower()
    assert "Q = 2 cm³/min" in r.text           # texto editável no SVG


def test_planilha_para_origin(cliente):
    r = cliente.post("/permeabilidade/origin", json=pedido())
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Permeabilidade"]
    assert [c.value for c in ws[1]][:2] == ["Tempo", "ΔP"]
    assert [c.value for c in ws[2]][:2] == ["s", "psi"]
    assert ws.max_column == 6                   # 3 vazões x (tempo, ΔP)
    assert wb["Etapas"].max_row == 2 + 3


def test_sem_login_e_recusado(monkeypatch):
    monkeypatch.delenv("MODO_DEV", raising=False)
    from app import app
    r = TestClient(app).post("/permeabilidade/grafico", json=pedido())
    assert r.status_code == 401


def test_nenhuma_etapa_visivel(cliente):
    p = pedido()
    for e in p["etapas"]:
        e["visivel"] = False
    assert cliente.post("/permeabilidade/grafico", json=p).status_code == 422


@pytest.mark.parametrize("sobrepor", [False, True])
def test_grafico_unico_e_o_padrao(cliente, sobrepor):
    p = pedido(formato="svg", sobrepor=sobrepor)
    p["config"].pop("modo", None)
    for i, cor in enumerate(["#000000", "#FF0000", "#0000FF"]):
        p["etapas"][i]["cor"] = cor
    r = cliente.post("/permeabilidade/grafico", json=p)
    assert r.status_code == 200
    svg = r.text.lower()
    assert svg.count("<g id=\"axes_") == 1                     # um gráfico só
    for cor in ("#000000", "#ff0000", "#0000ff"):
        assert cor in svg
    assert all(f"q = {q} cm³/min" in svg for q in (2, 4, 6))


def test_planilha_etapas_no_formato_do_pesquisador(cliente):
    etapas = [{**{k: e[k] for k in ("ordem", "q", "j_inicio", "j_fim", "dp_medio", "k_md", "usada")}, "n_pontos": 75}
              for e in PEDIDO["etapas"]]
    corpo = {"etapas": etapas, "comprimento_cm": 7.32, "diametro_cm": 3.79, "viscosidade_cp": 0.89}
    with open(AQUI / "permeabilidade_ficticio.xlsx", "rb") as f:
        r = cliente.post("/permeabilidade/planilha-etapas", files={"arquivo": ("ensaio.xlsx", f)}, data={"pedido": json.dumps(corpo)})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Planilha1"]
    assert list(ws.tables.keys()) == ["Tabela2"] and ws.tables["Tabela2"].ref.startswith("B2:")
    cab = [c.value for c in ws[2]][1:]
    assert cab == ["Tempo (s)", "PAH1 (psi)", "PAH2 (psi)", "dPH1 (psi)", "dPH1 (P2H1-P1H1) (psi)", "dPH2 (psi)",
                   "dPH2 (P2H2-P1H2) (psi)", "PC (psi)", "PBH1 (psi)", "PBH2 (psi)", "T. Estufa (oC)", "Q (cm³/min)",
                   "Nome G2", "Volume atual G2 (cm³)", "Volume bomb. parcial G2 (cm³)", "Nome G3", "Volume atual G3 (cm³)",
                   "Volume bomb. parcial G3 (cm³)", "Permeabilidade H1 (mD)", "Seleção do holder para PVBT", "Etapa"]
    rotulos = [ws.cell(row=i, column=len(cab) + 1).value for i in range(3, ws.max_row + 1)]
    assert sorted(set(rotulos)) == ["Permeabilidade 1", "Permeabilidade 2", "Permeabilidade 3"]
    assert ws.max_row - 2 == 225                         # 75 leituras estáveis por vazão
    assert wb["Resumo"]["B2"].value == "Permeabilidade por etapas de vazão"
    # Resumo traz todas as leituras extraídas de cada vazão: Tempo, ΔP e Volume, lado a lado
    rs = wb["Resumo"]
    titulos = [c.value for row in rs.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("Q = ")]
    assert titulos == ["Q = 2 cm³/min (Permeabilidade 1)", "Q = 4 cm³/min (Permeabilidade 2)", "Q = 6 cm³/min (Permeabilidade 3)"]
    lin_cab = next(c.row for row in rs.iter_rows() for c in row if c.value == "Tempo (s)")
    for c0 in (2, 6, 10):
        assert [rs.cell(row=lin_cab, column=c0 + k).value for k in range(3)] == ["Tempo (s)", "ΔP (psi)", "Volume (cm³)"]
        col = [rs.cell(row=i, column=c0).value for i in range(lin_cab + 1, rs.max_row + 1)]
        assert sum(v is not None for v in col) == 75
        vol = [rs.cell(row=i, column=c0 + 2).value for i in range(lin_cab + 1, lin_cab + 76)]
        assert vol == sorted(vol) and vol[-1] > vol[0]   # volume bombeado acumulado cresce


def test_planilha_etapas_com_graficos(cliente):
    etapas = [{k: e[k] for k in ("ordem", "q", "j_inicio", "j_fim", "dp_medio", "k_md", "usada")} for e in PEDIDO["etapas"]]
    corpo = {"etapas": etapas, "comprimento_cm": 7.32, "diametro_cm": 3.79, "viscosidade_cp": 0.89, "grafico": pedido()}
    with open(AQUI / "permeabilidade_ficticio.xlsx", "rb") as f:
        r = cliente.post("/permeabilidade/planilha-etapas", files={"arquivo": ("ensaio.xlsx", f)}, data={"pedido": json.dumps(corpo)})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Planilha1", "Resumo", "Gráficos"]
    ws = wb["Gráficos"]
    assert len(ws._images) == 2                          # gráfico único e ΔP × Q
    assert [ws.cell(row=c.row, column=2).value for c in ws["B"] if c.value in ("ΔP × tempo", "ΔP × Q")] == ["ΔP × tempo", "ΔP × Q"]


def test_planilha_etapas_sem_grafico_nao_tem_aba(cliente):
    etapas = [{k: e[k] for k in ("ordem", "q", "j_inicio", "j_fim", "dp_medio", "k_md", "usada")} for e in PEDIDO["etapas"]]
    with open(AQUI / "permeabilidade_ficticio.xlsx", "rb") as f:
        r = cliente.post("/permeabilidade/planilha-etapas", files={"arquivo": ("ensaio.xlsx", f)}, data={"pedido": json.dumps({"etapas": etapas})})
    assert r.status_code == 200, r.text
    assert "Gráficos" not in load_workbook(io.BytesIO(r.content)).sheetnames


def test_ajuste_igual_ao_origin():
    """O k final do laboratório usa a inclinação da reta com intercepto (FitLinear do Origin). Valores fictícios."""
    from graficos.planilha_etapas import ajuste_darcy
    a, b, r2, *_ = ajuste_darcy([(2, 30.0), (4, 57.0), (6, 84.0)])
    assert abs(a - 13.5) < 1e-9 and abs(b - 3) < 1e-9 and abs(r2 - 1) < 1e-12
    a0, b0, *_ = ajuste_darcy([(2, 30.0), (4, 57.0), (6, 84.0)], "origem")
    assert b0 == 0 and abs(a0 - 14.1428571) < 1e-6


# ------------------------------------------------------------------ acidificação (dados fictícios)
PEDIDO_ACID = {"amostra": "TESTE-A", "comprimento_cm": 7.3, "diametro_cm": 3.8, "porosidade_pct": 14,
               "volume_morto_cm3": 0.5, "temperatura_rotulo": "T40"}


def _post_acid(cliente, rota, pedido=None, arquivo: bytes | None = None, **form):
    conteudo = arquivo if arquivo is not None else (AQUI / "acidificacao_ficticio.xlsx").read_bytes()
    return cliente.post(rota, files={"arquivo": ("ensaio.xls", conteudo)}, data={"pedido": json.dumps(pedido or PEDIDO_ACID), **form})


def test_acid_detecta_fases_sem_ajuda():
    from graficos.acidificacao import Dados, detectar_fases
    d = Dados((AQUI / "acidificacao_ficticio.xlsx").read_bytes())
    fs = detectar_fases(d)
    assert [(f.fluido, f.q, f.papel) for f in fs] == [("KCl 2%", 2, "pre_fluxo"), ("BLEND", 2, "acido"),
                                                     ("AUP", 2, "pos_fluxo"), ("AUP", 4, "pos_fluxo"), ("AUP", 6, "pos_fluxo")]


def test_acid_planilha_igual_ao_origin(cliente):
    r = _post_acid(cliente, "/acidificacao/planilha")
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["0- Initial Data", "1 - Water Injection - Before ac", "2 - Acidizing",
                             "3 - Experiment general View", "4 - Temperature", "5 - Water Injection - After aci", "Gráficos"]
    ini = wb["0- Initial Data"]
    assert ini["B8"].value == "=PI()*((B4^2)/4)*B5*B6" and ini["B10"].value == "=B3/(B9*B6)"
    ac = wb["2 - Acidizing"]
    assert [ac.cell(row=1, column=c).value for c in range(2, 7)] == ["Time", "Volume injetado Total", "Injected Volume", "Pore Volume (PV)", "DP"]
    assert ac["C5"].value.endswith("*(B5-$B$5)/60") and ac["D5"].value.endswith("!$B$7") and ac["E5"].value.endswith("!$B$8")
    assert ac["F3"].value == "BLEND @2 ml/min T40"
    depois = wb["5 - Water Injection - After aci"]
    assert [depois.cell(row=3, column=c).value for c in (3, 5, 7)] == ["AUP 2.0 ml/min (after)", "AUP 4.0 ml/min (after)", "AUP 6.0 ml/min (after)"]
    assert len(wb["Gráficos"]._images) == 3


def test_acid_pv_igual_formula_do_origin():
    import math
    from graficos.acidificacao import Dados, PedidoAcid, detectar_fases, serie_acido
    d = Dados((AQUI / "acidificacao_ficticio.xlsx").read_bytes())
    p = PedidoAcid(**PEDIDO_ACID)
    fa = next(f for f in detectar_fases(d) if f.papel == "acido")
    s = serie_acido(d, fa, p)
    vp = math.pi * 3.8 ** 2 / 4 * 7.3 * 0.14
    t0 = s[0][0]
    for t, vt, vi, pv, _ in s[::10]:
        assert abs(vt - 2 * (t - t0) / 60) < 1e-12 and abs(vi - (vt - 0.5)) < 1e-12 and abs(pv - vi / vp) < 1e-12
    assert abs(s[0][3] + 0.5 / vp) < 1e-12


@pytest.mark.parametrize("qual", ["geral", "pv", "apos"])
def test_acid_graficos(cliente, qual):
    r = _post_acid(cliente, "/acidificacao/grafico", qual=qual)
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"\x89PNG"


# ------------------------------------------------------------------ formatos do arquivo do equipamento
# o mesmo ensaio fictício salvo como .xls antigo (BIFF8), HTML, XML 2003 e CSV com ";" e vírgula decimal,
# todos com extensão .xls: o serviço reconhece pelo conteúdo e dá o mesmo resultado do .xlsx
@pytest.mark.parametrize("nome", ["biff8.xls", "html.xls", "xml2003.xls", "csv_pt.xls"])
def test_acid_le_outros_formatos(cliente, nome):
    from graficos.acidificacao import Dados, detectar_fases
    ref = [(f.fluido, f.q, f.papel) for f in detectar_fases(Dados((AQUI / "acidificacao_ficticio.xlsx").read_bytes()))]
    conteudo = (AQUI / "formatos" / nome).read_bytes()
    assert [(f.fluido, f.q, f.papel) for f in detectar_fases(Dados(conteudo))] == ref
    r = _post_acid(cliente, "/acidificacao/planilha", arquivo=conteudo)
    assert r.status_code == 200, r.text
    assert len(load_workbook(io.BytesIO(r.content)).sheetnames) == 7


def test_arquivo_que_nao_e_planilha_da_422(cliente):
    r = _post_acid(cliente, "/acidificacao/planilha", arquivo=b"isto nao e uma planilha")
    assert r.status_code == 422, r.text


def _copia_fict(numeros_como_texto=False, aba_antes=False) -> bytes:
    """O ensaio fictício com números gravados como texto ("12,5") e/ou outra aba de leituras antes da bruta."""
    wb = load_workbook(AQUI / "acidificacao_ficticio.xlsx")
    ws = wb.worksheets[0]
    if numeros_como_texto:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.value = str(c.value).replace(".", ",")
    if aba_antes:
        outra = wb.create_sheet("Separada", 0)          # tabela já separada, com tempos que não batem com a bruta
        outra.append(["Tempo (s)", "Q (cm³/min)", "dPH1 (psi)"])
        for t in range(0, 50, 5):
            outra.append([t, 2, 10])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("texto,antes", [(True, False), (False, True), (True, True)])
def test_acid_mesmas_leituras_que_o_gerencial(cliente, texto, antes):
    from graficos.acidificacao import Dados, detectar_fases
    ref = Dados((AQUI / "acidificacao_ficticio.xlsx").read_bytes())
    fases = detectar_fases(ref)
    conteudo = _copia_fict(texto, antes)
    d = Dados(conteudo, ref.aba)                        # a aba que o gerencial manda
    assert len(d.linhas) == len(ref.linhas)
    pedido = {**PEDIDO_ACID, "aba": ref.aba, "fases": [f.model_dump() for f in fases]}
    r = _post_acid(cliente, "/acidificacao/planilha", pedido=pedido, arquivo=conteudo)
    assert r.status_code == 200, r.text
    # sem o nome da aba, prefere a exportação bruta (com o cabeçalho), como o gerencial
    assert Dados(conteudo).aba == ref.aba


def test_acid_fase_fora_do_arquivo_explica(cliente):
    pedido = {**PEDIDO_ACID, "fases": [{"ordem": 1, "papel": "acido", "q": 2, "t_inicio": 99999, "t_fim": 99999.5}]}
    r = _post_acid(cliente, "/acidificacao/planilha", pedido=pedido)
    assert r.status_code == 422 and "aba" in r.text and "leituras" in r.text, r.text
