"""
Gera uma planilha FICTÍCIA com a mesma estrutura da BDD_LCPETRO_Reatividade
(mesmas abas e colunas), para desenvolver e testar o importador sem usar os
dados reais do laboratório.

Os valores são inventados. Alguns erros foram colocados DE PROPÓSITO para ver
se o importador os pega (ver ERROS_PROPOSITAIS no fim do arquivo).

Uso:  python gerar_planilha_ficticia.py
Saída: BDD_ficticio.xlsx nesta pasta
"""
import math
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

random.seed(42)
SAIDA = Path(__file__).with_name("BDD_ficticio.xlsx")

AZUL = PatternFill("solid", fgColor="1F3864")
VERDE = PatternFill("solid", fgColor="2E7D4F")
BRANCO = Font(color="FFFFFF", bold=True)

# ----------------------------------------------------------------------------
referencias = [
    ("REF-R01", "Autor Ficticio, A. et al", "2019"),
    ("REF-R02", "Pesquisador, B. et al", "2010"),
    ("REF-R03", "Exemplo, C. et al", "2006-2014"),   # faixa de anos, como na real
    ("REF-R04", "Teste, D. at al", "1973"),          # "at al" de propósito
]

litologias = ["Indiana Limestone", "Austin Chalk", "Alabama Marble",
              "Dolomita", "Pink Desert Limestone", "Calcita"]

rochas = [
    # id, litologia, composicao_texto, calcita, dolomita, razao
    ("RCH-01", "Indiana Limestone", "[Calcita] = 98%", 98, None, None),
    ("RCH-02", "Indiana", "[Calcita] > 97 %", 97, None, None),      # sinônimo
    ("RCH-03", "Austin Chalk", "[Calcita] > 99 %", 99, None, None),
    ("RCH-04", "Alamaba Marble", "[Calcita] = 96%", 96, None, None), # erro de grafia
    ("RCH-05", "Dolomita", "RM Ca/Mg = 1,1", None, None, 1.1),
    ("RCH-06", "Dolomita", "RM Ca/Mg = 1,4", None, None, 1.4),
    ("RCH-07", "Pink Desert Limestone", "", None, None, None),
    ("RCH-08", "Calcita", "[Calcita]= 99%", 99, None, None),
]

fluidos = [
    ("HCl", 36.46, 0.00494, 1, "ρ = 1 + 0,00494·wt% (20 °C); 1 N = 1 M"),
    ("Ácido cítrico", 192.12, 0.0035, 3, "triprótico: 1 N = 1/3 M"),
    ("Ácido acético", 60.05, 0.00125, 1, ""),
    ("GLDA", None, None, None, "quelante: conversão depende da forma salina"),
]

sistemas = [
    # id, nome, [(comp, valor, unidade) ...], original
    ("SA01", "HCl", [("HCl", 0.1, "N")], "[HCl] = 0,1 N"),
    ("SA02", "HCl", [("HCl", 0.5, "M")], "[HCl] = 0,5 M"),
    ("SA03", "HCl 15%", [("HCl", 15, "wt%")], "[HCl] = 15 wt%"),
    ("SA04", "HCl + Surfact", [("HCl", 0.1, "N"), ("Surfact", 0.2, "vol%")],
     "[HCl] = 0,1 N [Surfact] = 0,2 vol%"),
    ("SA05", "Ácido cítrico", [("Ácido cítrico", 0.3, "N")], "[Cítrico] = 0,3 N"),
    ("SA06", "Ácido acético", [("Ácido acético", 10, "wt%")], "[HAc] = 10 wt%"),
]

# ----------------------------------------------------------------------------
# ensaios: série tipo Levich (taxa ~ k·√ω) com ruído
COLS_ENSAIOS = [
    "id_ensaio", "id_referencia", "ano", "id_rocha", "litologia", "composicao_rocha",
    "calcita_pct", "dim_h_in", "dim_d_in", "id_sistema", "acido", "conc_valor",
    "conc_unidade", "conc_pct", "conc_M", "aditivos", "temperatura_C",
    "temperatura_nominal_C", "pressao_psi", "tipo_variavel", "rotacao_rpm",
    "omega_rad_s", "raiz_omega", "shear_rate_s1", "taxa_reacao_mol_cm2_s",
    "taxa_reacao_mol_m2_s", "coef_difusao_cm2_s", "metodo_coef_difusao",
    "metodo_quantificacao", "obtencao_dado", "cor_celula", "variavel_original",
    "sistema_original",
]
CALCULADAS = {"litologia", "composicao_rocha", "calcita_pct", "acido", "conc_valor",
              "conc_unidade", "conc_pct", "conc_M", "aditivos", "temperatura_nominal_C",
              "omega_rad_s", "raiz_omega", "taxa_reacao_mol_m2_s"}

ensaios = []
n = 0
plano = [
    ("REF-R01", 2019, "RCH-01", "SA01", 23.0, [100, 200, 400, 800, 1200], 6.0e-7, "tabelado"),
    ("REF-R01", 2019, "RCH-03", "SA01", 23.0, [100, 200, 400, 800], 5.5e-7, "tabelado"),
    ("REF-R02", 2010, "RCH-05", "SA02", 50.0, [100, 300, 500, 1000], 2.0e-7, "lido de gráfico"),
    ("REF-R02", 2010, "RCH-06", "SA02", 50.0, [100, 300, 500, 1000], 2.6e-7, "lido de gráfico"),
    ("REF-R03", 2008, "RCH-02", "SA04", 93.3, [100, 500, 1000], 1.1e-6, "tabelado"),
    ("REF-R03", 2012, "RCH-04", "SA05", 121.1, [200, 600], 3.0e-7, "medida na superfície"),
    ("REF-R04", 1973, "RCH-08", "SA03", 25.0, [50, 150, 450, 900], 4.0e-6, "não catalogado"),
    ("REF-R04", 1973, "RCH-07", "SA06", 148.8, [100, 400], 8.0e-8, "tabelado"),
]
for ref, ano, rocha, sist, temp, rpms, k, obt in plano:
    for rpm in rpms:
        n += 1
        w = rpm * 2 * math.pi / 60
        taxa = k * math.sqrt(w) / 3.236 * random.uniform(0.9, 1.1)
        ensaios.append({
            "id_ensaio": f"RDA-{n:04d}", "id_referencia": ref, "ano": ano,
            "id_rocha": rocha, "dim_h_in": 1.5, "dim_d_in": 1.0, "id_sistema": sist,
            "temperatura_C": temp, "pressao_psi": 1000, "tipo_variavel": "rotação (RDA)",
            "rotacao_rpm": rpm, "taxa_reacao_mol_cm2_s": float(f"{taxa:.3e}"),
            "metodo_quantificacao": "RDA / ICP-AES", "obtencao_dado": obt,
            "cor_celula": "FFFFFF", "variavel_original": f"{rpm} rpm",
            "sistema_original": sist,
        })

# linha de conferência dos cálculos: 0,1 N HCl, 100 rpm, 23 °C, taxa conhecida
ensaios[0]["taxa_reacao_mol_cm2_s"] = 3.10e-7

# um ensaio por taxa de cisalhamento, sem rotação
n += 1
ensaios.append({"id_ensaio": f"RDA-{n:04d}", "id_referencia": "REF-R02", "ano": 2010,
                "id_rocha": "RCH-05", "id_sistema": "SA02", "temperatura_C": 50,
                "pressao_psi": "-", "tipo_variavel": "taxa de cisalhamento",
                "shear_rate_s1": 250, "taxa_reacao_mol_cm2_s": 2.2e-7,
                "metodo_quantificacao": "RDA/ ICP", "obtencao_dado": "tabelado",
                "variavel_original": "250 1/s", "sistema_original": "HCl 0,5 M"})

# ---- ERROS_PROPOSITAIS -------------------------------------------------------
# 1. rocha que não existe no catálogo
ensaios[3]["id_rocha"] = "RCH-99"
# 2. valor não numérico na taxa
ensaios[5]["taxa_reacao_mol_cm2_s"] = "5,1e-7?"
# 3. id duplicado
ensaios[7]["id_ensaio"] = ensaios[6]["id_ensaio"]
# 4. taxa fora de faixa plausível (provável erro de unidade)
ensaios[9]["taxa_reacao_mol_cm2_s"] = 3.2e-3
# 5. valor fora da lista
ensaios[11]["obtencao_dado"] = "Tabelado "   # maiúscula + espaço -> deve ser corrigido
ensaios[12]["tipo_variavel"] = "rpm"         # não existe na lista -> erro


def cabecalho(ws, cols, verdes=()):
    ws.append(cols)
    for c, nome in zip(ws[1], cols):
        c.fill = VERDE if nome in verdes else AZUL
        c.font = BRANCO


wb = Workbook()
ws = wb.active
ws.title = "LEIA-ME"
ws["A1"] = "PLANILHA FICTÍCIA — mesma estrutura da BDD_LCPETRO_Reatividade, valores inventados"

ws = wb.create_sheet("dicionario")
cabecalho(ws, ["tabela", "campo", "tipo", "unidade", "valores_aceitos", "descricao"])

ws = wb.create_sheet("ensaios")
cabecalho(ws, COLS_ENSAIOS, CALCULADAS)
for e in ensaios:
    ws.append([e.get(c) for c in COLS_ENSAIOS])

ws = wb.create_sheet("rochas")
cabecalho(ws, ["id_rocha", "litologia", "composicao_texto", "calcita_pct",
               "dolomita_pct", "razao_ca_mg"])
for r in rochas:
    ws.append(list(r))

ws = wb.create_sheet("sistemas_acidos")
cols = ["id_sistema", "nome_sistema"]
for i in range(1, 5):
    cols += [f"componente_{i}", f"valor_{i}", f"unidade_{i}"]
cols.append("composicao_original")
cabecalho(ws, cols)
for sid, nome, comps, orig in sistemas:
    linha = [sid, nome]
    for i in range(4):
        linha += list(comps[i]) if i < len(comps) else [None, None, None]
    ws.append(linha + [orig])

ws = wb.create_sheet("fluidos_prop")
cabecalho(ws, ["componente", "massa_molar_g_mol", "coef_densidade", "n_equivalentes",
               "observacao"])
for f in fluidos:
    ws.append(list(f))

ws = wb.create_sheet("referencias")
cabecalho(ws, ["id_referencia", "autor", "ano", "n_ensaios"])
for rid, autor, ano in referencias:
    ws.append([rid, autor, ano, 99])   # n_ensaios digitado errado de propósito

ws = wb.create_sheet("qualidade")
cabecalho(ws, ["id_ensaio", "problema"])

ws = wb.create_sheet("listas")
listas = {
    "tipo_variavel": ["rotação (RDA)", "taxa de cisalhamento", "não informada"],
    "obtencao_dado": ["tabelado", "lido de gráfico", "medida na superfície", "não catalogado"],
    "conc_unidade": ["N", "M", "wt%"],
    "metodo_quantificacao": ["RDA", "RDA / ICP", "RDA / ICP-AES", "RDA/ Coreflooding"],
    "metodo_coef_difusao": ["Newman, J. (1966)", "Rozieres et al (1994)"],
    "litologia": litologias,
}
cabecalho(ws, list(listas))
for i in range(max(map(len, listas.values()))):
    ws.append([v[i] if i < len(v) else None for v in listas.values()])

wb.save(SAIDA)
print(f"{SAIDA.name}: {len(ensaios)} ensaios, {len(rochas)} rochas, "
      f"{len(sistemas)} sistemas, {len(referencias)} referências")
