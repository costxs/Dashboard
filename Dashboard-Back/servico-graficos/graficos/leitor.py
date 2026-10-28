"""Leitura do arquivo exportado pelos dois equipamentos de coreflood, num formato único para o serviço.

JBV  cabeçalho "Chave: valor" + tabela com Tempo (s), dPH1/dPH2 (psi), Q (cm³/min), garrafas G1..G4.
DCI  log .tsv do software (aberto no Excel): cabeçalho "Chave | Valor" (Sample ID, Diam/Length em polegadas,
     Porosity, Pore Volume, viscosidade de cada acumulador) + tabela com Timestamp (hours), VPA n (bombas),
     Core B DP1B/DP2B/DP3B e DPBtot (psi), Accum. 1..4 Vol (ml). Os números vêm como texto (ponto decimal).
     Também aceita a cópia que o pesquisador faz numa aba (tabela começando na coluna C, sem cabeçalho).

Tabela.t(l) devolve o tempo em SEGUNDOS nos dois casos (no DCI, horas × 3600), igual ao gerencial.
"""
from __future__ import annotations

import io
import re
import unicodedata
from typing import Optional

from openpyxl import load_workbook


def chave(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _num(v):
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return v
    return v


FIXAS_JBV = ["tempo s", "pah1 psi", "pah2 psi", "dph1 psi", "dph1 p2h1 p1h1 psi", "dph2 psi", "dph2 p2h2 p1h2 psi",
             "pc psi", "pbh1 psi", "pbh2 psi", "t estufa oc", "q cm min"]


class Tabela:
    equipamento: str
    meta: dict              # amostra, data, operador, comprimento_cm, diametro_cm, porosidade_pct, viscosidade_cp, fluido, volume_poroso_cm3
    meta_bruta: dict
    cab: list
    linhas: list
    idx: dict
    garrafas: dict          # n -> (coluna do volume atual, coluna do nome ou -1)
    nome_garrafa: str       # "garrafa" (JBV) ou "acumulador" (DCI)
    rotulo_dp: str
    escala_t: float = 1.0

    def t(self, l) -> float:
        return l[self.idx["t"]] * self.escala_t

    def v(self, l, nome: str) -> Optional[float]:
        j = self.idx.get(nome, -1)
        x = l[j] if j is not None and j >= 0 else None
        return x if isinstance(x, (int, float)) else None

    def colunas_planilha(self) -> list[int]:
        return self._cols


def ler_arquivo(conteudo: bytes) -> Tabela:
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    abas = [(ws.title, list(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
    # DCI primeiro (tem "Timestamp (hours)" e "VPA n Flow Rate")
    dci = []
    for titulo, linhas in abas:
        for i, l in enumerate(linhas[:80]):
            ks = [chave(v) for v in l]
            if "timestamp hours" in ks and any(re.match(r"^vpa \d+ flow rate", k) for k in ks):
                dci.append((titulo, linhas, i))
                break
    if dci:
        return _dci(dci)
    for titulo, linhas in abas:
        for i, l in enumerate(linhas):
            ks = [chave(v) for v in l]
            if "tempo s" in ks and any(k.startswith("q cm") for k in ks):
                return _jbv(linhas, i)
    raise ValueError("não reconheci o arquivo: esperava a exportação do JBV (Tempo (s), Q (cm³/min)) "
                     "ou do DCI (Timestamp (hours), VPA n Flow Rate)")


# --------------------------------------------------------------------------- JBV
def _jbv(linhas, i) -> Tabela:
    tb = Tabela()
    tb.equipamento, tb.nome_garrafa = "JBV", "garrafa"
    meta = {}
    for m in linhas[:i]:
        t = next((v for v in m if isinstance(v, str) and ":" in v), None)
        if t:
            a, b = t.split(":", 1)
            meta[a.strip()] = b.strip()
    tb.meta_bruta = meta
    num = lambda s: float(s.replace(",", ".")) if s and re.match(r"^-?[\d.,]+$", s) else None  # noqa: E731
    tb.meta = {"amostra": meta.get("Descrição da amostra"), "data": meta.get("Data"), "operador": meta.get("Usuário"),
               "comprimento_cm": num(meta.get("Comprimento (cm)")), "diametro_cm": num(meta.get("Diâmetro (cm)")),
               "porosidade_pct": num(meta.get("Porosidade (%)")), "viscosidade_cp": num(meta.get("Viscosidade (cP)")),
               "fluido": meta.get("Fluido de tratamento") or None, "volume_poroso_cm3": None}
    tb.cab = list(linhas[i])
    ks = [chave(v) for v in tb.cab]
    i_t = ks.index("tempo s")
    tb.linhas = [l for l in linhas[i + 1:] if l and isinstance(l[i_t], (int, float))]
    col = lambda f: next((j for j, k in enumerate(ks) if f(k)), -1)  # noqa: E731
    sel = col(lambda k: k.startswith("selecao do holder"))
    h2 = sel >= 0 and sum("2" in str(l[sel] or "") for l in tb.linhas) > sum("1" in str(l[sel] or "") for l in tb.linhas)
    h = "h2" if h2 else "h1"
    tb.idx = {"t": i_t, "q": col(lambda k: k.startswith("q cm")), "dp": col(lambda k: k == f"dp{h} psi"),
              "pb": col(lambda k: k == f"pb{h} psi"), "pc": col(lambda k: k == "pc psi"),
              "temp": col(lambda k: k.startswith("t estufa")), "k": col(lambda k: k == f"permeabilidade {h} md"),
              "vol_bomba": -1}
    tb.rotulo_dp = f"dP{h.upper()}"
    if tb.idx["dp"] < 0 or tb.idx["q"] < 0:
        raise ValueError(f"colunas dP{h.upper()} (psi) e Q (cm³/min) não encontradas")
    tb.garrafas = {}
    for j, k in enumerate(ks):
        m = re.match(r"^volume atual g(\d)", k)
        if m:
            n = int(m.group(1))
            tb.garrafas[n] = (j, col(lambda x, n=n: x == f"nome g{n}"))
    # colunas da Planilha1 (iguais às que o pesquisador mantém)
    cols = [ks.index(f) for f in FIXAS_JBV if f in ks]
    for g in range(1, 9):
        i_nome = ks.index(f"nome g{g}") if f"nome g{g}" in ks else -1
        if i_nome >= 0 and any(l[i_nome] not in (None, "") for l in tb.linhas):
            cols += [i_nome] + [j for j, k in enumerate(ks) if k in (f"volume atual g{g} cm", f"volume bomb parcial g{g} cm")]
    if tb.idx["k"] >= 0:
        cols.append(tb.idx["k"])
    if sel >= 0:
        cols.append(sel)
    tb._cols = cols
    return tb


# --------------------------------------------------------------------------- DCI
def _dci(achados) -> Tabela:
    # cabeçalho: de qualquer aba que o tenha
    kv = {}
    for _, linhas, i in achados:
        for l in linhas[:i]:
            vals = [v for v in l if v not in (None, "")]
            if len(vals) >= 2 and isinstance(vals[0], str):
                kv.setdefault(chave(vals[0]), str(vals[1]).strip())
            elif len(vals) == 1 and isinstance(vals[0], str):
                kv.setdefault(chave(vals[0]), "")
        if "log start time" in kv:
            break
    # dados: prefere a aba com a separação do pesquisador (coluna com "Permeabilidade n" ou "... @Q")
    def tem_rotulos(linhas, i):
        return any(isinstance(v, str) and re.search(r"permeab|@", v, re.I) for l in linhas[i + 1:i + 400] for v in l)
    achados = sorted(achados, key=lambda a: not tem_rotulos(a[1], a[2]))
    titulo, linhas, i = achados[0]

    tb = Tabela()
    tb.equipamento, tb.nome_garrafa, tb.escala_t = "DCI", "acumulador", 3600.0
    f = lambda k: _num(kv.get(k))  # noqa: E731
    temB = bool(kv.get("sample id core b")) or (isinstance(f("diam core b in"), float) and f("diam core b in") > 0.05)
    c = "b" if temB else "a"
    pol = lambda k: (f(k) * 2.54) if isinstance(f(k), float) and f(k) > 0.05 else None  # noqa: E731
    visc = {g: f(f"accum {g} viscosity cp") for g in range(1, 5) if isinstance(f(f"accum {g} viscosity cp"), float)}
    proj = kv.get("project name") or ""
    mdata = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})\D+(\d{1,2}):(\d{2}):(\d{2})\s*(AM|PM)?", kv.get("log start time", ""), re.I)
    data = None
    if mdata:
        hh = int(mdata.group(4)) + (12 if (mdata.group(7) or "").upper() == "PM" and int(mdata.group(4)) < 12 else 0)
        hh = 0 if (mdata.group(7) or "").upper() == "AM" and int(mdata.group(4)) == 12 else hh
        data = f"{int(mdata.group(2)):02d}/{int(mdata.group(1)):02d}/{mdata.group(3)} {hh:02d}:{mdata.group(5)}:{mdata.group(6)}"
    partes = re.sub(r"@\s*[\d.,]+|\bT\s*\d+(?:[.,]\d+)?\b", "", proj).split()
    tb.meta_bruta = kv
    tb.meta = {"amostra": kv.get(f"sample id core {c}") or (partes[0] if partes else None), "data": data,
               "operador": kv.get("operator"), "comprimento_cm": pol(f"length core {c} in"), "diametro_cm": pol(f"diam core {c} in"),
               "porosidade_pct": f(f"porosity core {c}") if isinstance(f(f"porosity core {c}"), float) and f(f"porosity core {c}") > 0 else None,
               "viscosidade_cp": next(iter(visc.values()), None), "fluido": " ".join(partes[1:]) or None,
               "volume_poroso_cm3": f(f"pore volume core {c} ml") if isinstance(f(f"pore volume core {c} ml"), float) and f(f"pore volume core {c} ml") > 0 else None,
               "projeto": proj}

    tb.cab = list(linhas[i])
    ks = [chave(v) for v in tb.cab]
    col = lambda g: next((j for j, k in enumerate(ks) if g(k)), -1)  # noqa: E731
    i_t = ks.index("timestamp hours")
    tb.linhas = []
    for l in linhas[i + 1:]:
        if not l:
            continue
        ln = [_num(v) for v in l]
        if isinstance(ln[i_t], (int, float)):
            tb.linhas.append(ln)
    # bomba de injeção = a VPA que mais bombeou
    bombas = []
    for j, k in enumerate(ks):
        m = re.match(r"^vpa (\d+) volume pumped ml$", k)
        if m:
            vs = [l[j] for l in tb.linhas if isinstance(l[j], (int, float))]
            bombas.append(((max(vs) - min(vs)) if vs else 0, m.group(1), j))
    bombas.sort(reverse=True)
    b = bombas[0][1] if bombas else "1"
    tb.idx = {"t": i_t, "q": col(lambda k: k == f"vpa {b} flow rate ml min"),
              "dp": col(lambda k: k.startswith(f"core {c} differential pressure total")),
              "pb": col(lambda k: k.startswith(f"core {c} downstream pressure")),
              "pc": col(lambda k: k.startswith(f"core {c} confining pressure")),
              "temp": col(lambda k: k.startswith("oven temp")), "k": col(lambda k: k == f"permeability 1{c} md"),
              "vol_bomba": bombas[0][2] if bombas else -1}
    tb.rotulo_dp = f"DP{c.upper()}tot"
    if tb.idx["dp"] < 0 or tb.idx["q"] < 0:
        raise ValueError("colunas Differential Pressure Total e VPA Flow Rate não encontradas no arquivo do DCI")
    tb.garrafas = {}
    for g in range(1, 5):
        j = col(lambda k, g=g: k == f"accum {g} vol ml")
        if j >= 0:
            tb.garrafas[g] = (j, -1)
    # colunas da Planilha1: as mesmas que o pesquisador copia (sem data/hora/bombas de pressão)
    manter = [i_t] + [col(lambda k, s=s: k == f"vpa {b} {s}") for s in ("volume pumped ml", "pressure psi", "flow rate ml min")]
    manter += [col(lambda k, s=s: k.startswith(s)) for s in ("oven temp", f"core {c} upstream pressure", f"core {c} downstream pressure",
                                                             f"core {c} differential pressure 1", f"core {c} differential pressure 2",
                                                             f"core {c} differential pressure 3", f"core {c} confining pressure",
                                                             f"core {c} fluid temp in", f"core {c} fluid temp out")]
    for g, (j, _) in tb.garrafas.items():
        vs = [l[j] for l in tb.linhas if isinstance(l[j], (int, float))]
        if vs and max(vs) - min(vs) > 0.01:
            manter.append(j)
    manter += [tb.idx["k"], tb.idx["dp"]]
    tb._cols = [j for j in manter if j is not None and j >= 0]
    return tb


# --------------------------------------------------------------------------- adaptador para o formato do JBV
def eh_dci(conteudo: bytes) -> bool:
    """O arquivo é a exportação do DCI (.xlsx)? Qualquer erro de leitura = não é (cai no leitor do JBV)."""
    try:
        wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
        for ws in wb.worksheets:
            for i, l in enumerate(ws.iter_rows(values_only=True)):
                if i >= 80:
                    break
                ks = [chave(v) for v in l]
                if "timestamp hours" in ks and any(re.match(r"^vpa \d+ flow rate", k) for k in ks):
                    return True
    except Exception:  # noqa: BLE001
        return False
    return False


class DCIComoJBV:
    """O arquivo do DCI com as mesmas colunas, nomes e unidades do JBV: tempo em s, dPH1 (psi), Q (cm³/min),
    Nome Gn / Volume atual Gn (cm³) para cada acumulador. Todo o resto do serviço (planilha de etapas,
    gráficos e planilha de acidificação) lê o JBV e passa a servir ao DCI sem mudança."""
    nome_aba = "DCI"

    def __init__(self, meta: dict, cab: list, linhas: list):
        self.meta, self.cab, self.linhas = meta, cab, linhas


def dci_como_jbv(conteudo: bytes, dp: str = "total") -> DCIComoJBV:
    tb = ler_arquivo(conteudo)
    if tb.equipamento != "DCI":
        raise ValueError("o arquivo não é do DCI")
    ks = [chave(v) for v in tb.cab]
    c = tb.rotulo_dp[2].lower()                      # "DPBtot" -> "b"
    j_dp1 = next((j for j, k in enumerate(ks) if k.startswith(f"core {c} differential pressure 1")), -1)
    j_dp2 = next((j for j, k in enumerate(ks) if k.startswith(f"core {c} differential pressure 2")), -1)
    soma = dp == "dp1_dp2" and j_dp1 >= 0 and j_dp2 >= 0

    def num(l, j):
        x = l[j] if 0 <= j < len(l) else None
        return x if isinstance(x, (int, float)) else None

    def dp_de(l):
        if soma:
            a, b = num(l, j_dp1), num(l, j_dp2)
            return (a + b) if a is not None and b is not None else None
        return num(l, tb.idx["dp"])

    # só os acumuladores que mudaram de volume (os parados não são fluido em uso)
    gs = []
    for g, (j, _) in tb.garrafas.items():
        vs = [l[j] for l in tb.linhas if isinstance(l[j], (int, float))]
        if vs and max(vs) - min(vs) > 0.01:
            gs.append((g, j))
    cab = ["Tempo (s)", "dPH1 (psi)", "Q (cm³/min)", "PC (psi)", "PBH1 (psi)", "T. Estufa (oC)", "Permeabilidade H1 (mD)"]
    for g, _ in gs:
        cab += [f"Nome G{g}", f"Volume atual G{g} (cm³)"]
    linhas = []
    for l in tb.linhas:
        linha = [tb.t(l), dp_de(l), num(l, tb.idx["q"]), num(l, tb.idx["pc"]), num(l, tb.idx["pb"]),
                 num(l, tb.idx["temp"]), num(l, tb.idx["k"])]
        for g, j in gs:
            linha += [f"Acumulador {g}", num(l, j)]
        linhas.append(linha)
    m = tb.meta
    def txt(v, casas=None):
        return None if v is None else (f"{v:.{casas}f}" if casas is not None else str(v))
    meta = {"Descrição da amostra": m.get("amostra"), "Data": (m.get("data") or "").split(" ")[0] or None, "Usuário": m.get("operador"),
            "Comprimento (cm)": txt(m.get("comprimento_cm"), 2), "Diâmetro (cm)": txt(m.get("diametro_cm"), 2),
            "Porosidade (%)": txt(m.get("porosidade_pct"), 2), "Viscosidade (cP)": txt(m.get("viscosidade_cp")),
            "Fluido de tratamento": m.get("fluido"), "Equipamento": "DCI"}
    return DCIComoJBV({k: v for k, v in meta.items() if v not in (None, "")}, cab, linhas)
