"""Planilha "separada por etapas" no formato que o pesquisador já usa (Planilha3 do exemplo):

  - aba "Planilha1": tabela do Excel (Tabela2) começando em B2, só com as colunas que importam
    para a permeabilidade, na mesma ordem do equipamento, e a última coluna dizendo a etapa
    ("Permeabilidade 1", "Permeabilidade 2"...). Só entram as leituras do TRECHO ESTÁVEL de cada vazão;
  - aba "Resumo": dados da amostra, uma linha por etapa (Q, trecho, ΔP médio ± desvio, volume, k)
    e o k final (ajuste ΔP × Q e média), formatada para imprimir.

Lê direto o arquivo exportado pelo equipamento (o mesmo que o gerencial leu), para não perder
nenhuma coluna nem casa decimal.
"""
from __future__ import annotations

import io
import math
import re
import unicodedata
from datetime import datetime
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pydantic import BaseModel

CONSTANTE_DARCY = 1000 * 14.696 / 60

# colunas mantidas, na ordem do equipamento (as demais — Hora, garrafas vazias, holder não usado,
# alarme, set point, PVBT, volumes produzidos — ficam de fora)
FIXAS = ["tempo s", "pah1 psi", "pah2 psi", "dph1 psi", "dph1 p2h1 p1h1 psi", "dph2 psi", "dph2 p2h2 p1h2 psi",
         "pc psi", "pbh1 psi", "pbh2 psi", "t estufa oc", "q cm min"]


def chave(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def ajuste_darcy(pts: list[tuple[float, float]], modo: str = "intercepto"):
    """Reta ΔP (psi) × Q (cm³/min). 'intercepto' = ΔP = a·Q + b (FitLinear do Origin, padrão do laboratório);
    'origem' = ΔP = a·Q. Devolve (a, b, r2, a_livre, b_livre, a_origem)."""
    n = len(pts)
    mq = sum(q for q, _ in pts) / n
    mp = sum(d for _, d in pts) / n
    sqq = sum((q - mq) ** 2 for q, _ in pts)
    a_l = sum((q - mq) * (d - mp) for q, d in pts) / sqq if sqq else float("nan")
    b_l = mp - a_l * mq
    a_0 = sum(q * d for q, d in pts) / sum(q * q for q, _ in pts)
    a, b = (a_0, 0.0) if modo == "origem" else (a_l, b_l)
    ss_tot = sum((d - mp) ** 2 for _, d in pts)
    r2 = 1 - sum((d - (a * q + b)) ** 2 for q, d in pts) / ss_tot if ss_tot else 1.0
    return a, b, r2, a_l, b_l, a_0


class EtapaJanela(BaseModel):
    ordem: int
    q: float
    j_inicio: float
    j_fim: float
    t_inicio: Optional[float] = None
    t_fim: Optional[float] = None
    dp_medio: Optional[float] = None
    dp_desvio: Optional[float] = None
    n_pontos: Optional[int] = None
    volume_injetado_cm3: Optional[float] = None
    volume_bombeado_cm3: Optional[float] = None
    k_md: Optional[float] = None
    usada: bool = True
    rotulo: Optional[str] = None


class PedidoPlanilha(BaseModel):
    etapas: list[EtapaJanela]
    metodo_k: str = "ajuste ΔP × Q"
    comprimento_cm: Optional[float] = None
    diametro_cm: Optional[float] = None
    viscosidade_cp: Optional[float] = None
    ajuste: str = "intercepto"   # 'intercepto' (como o Origin do laboratório) ou 'origem'


def _ler_exportacao(conteudo: bytes):
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        for i, l in enumerate(linhas):
            ks = [chave(v) for v in l]
            if "tempo s" in ks and any(k.startswith("q cm") for k in ks):
                meta = {}
                for m in linhas[:i]:
                    t = next((v for v in m if isinstance(v, str) and ":" in v), None)
                    if t:
                        a, b = t.split(":", 1)
                        meta[a.strip()] = b.strip()
                return meta, list(l), linhas[i + 1:]
    raise ValueError("não achei a tabela de leituras (colunas 'Tempo (s)' e 'Q (cm³/min)') no arquivo")


def _colunas_mantidas(cab: list, dados: list) -> list[int]:
    ks = [chave(v) for v in cab]
    idx = [ks.index(f) for f in FIXAS if f in ks]
    # garrafas que têm nome em alguma linha
    for g in range(1, 9):
        i_nome = ks.index(f"nome g{g}") if f"nome g{g}" in ks else -1
        if i_nome < 0:
            continue
        if any(l[i_nome] not in (None, "") for l in dados):
            idx += [i_nome] + [j for j, k in enumerate(ks) if k in (f"volume atual g{g} cm", f"volume bomb parcial g{g} cm")]
    # holder escolhido e a permeabilidade dele
    i_sel = next((j for j, k in enumerate(ks) if k.startswith("selecao do holder")), -1)
    h2 = i_sel >= 0 and sum("2" in str(l[i_sel] or "") for l in dados) > sum("1" in str(l[i_sel] or "") for l in dados)
    k_alvo = "permeabilidade h2 md" if h2 else "permeabilidade h1 md"
    if k_alvo in ks:
        idx.append(ks.index(k_alvo))
    if i_sel >= 0:
        idx.append(i_sel)
    return idx


AZUL = "1F4E9C"
FINO = Side(style="thin", color="B7C3D7")


def gerar_planilha_etapas(conteudo: bytes, p: PedidoPlanilha) -> bytes:
    meta, cab, dados = _ler_exportacao(conteudo)
    ks = [chave(v) for v in cab]
    i_t = ks.index("tempo s")
    dados = [l for l in dados if l and isinstance(l[i_t], (int, float))]
    cols = _colunas_mantidas(cab, dados)

    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    titulos = [str(cab[j]).strip() for j in cols] + ["Etapa"]
    for c, t in enumerate(titulos, start=2):
        ws.cell(row=2, column=c, value=t)
    linha = 3
    for e in sorted(p.etapas, key=lambda x: x.ordem):
        rot = e.rotulo or f"Permeabilidade {e.ordem}"
        for l in dados:
            if e.j_inicio <= l[i_t] <= e.j_fim:
                for c, j in enumerate(cols, start=2):
                    ws.cell(row=linha, column=c, value=l[j])
                ws.cell(row=linha, column=len(cols) + 2, value=rot)
                linha += 1
    if linha == 3:
        raise ValueError("nenhuma leitura dentro dos trechos estáveis informados")
    ult = get_column_letter(len(titulos) + 1)
    tab = Table(displayName="Tabela2", ref=f"B2:{ult}{linha - 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    for c, t in enumerate(titulos, start=2):
        letra = get_column_letter(c)
        ws.column_dimensions[letra].width = max(10, min(28, len(t) + 2))
        ws.cell(row=2, column=c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for r in range(3, linha):
            cel = ws.cell(row=r, column=c)
            if isinstance(cel.value, float):
                cel.number_format = "0.00"
        cab_cel = ws.cell(row=2, column=c)
        cab_cel.font = Font(bold=True, color="FFFFFF")
        cab_cel.fill = PatternFill("solid", fgColor=AZUL)
    ws.column_dimensions["A"].width = 3
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "C3"
    _impressao(ws)

    # ---------------- Resumo ----------------
    r = wb.create_sheet("Resumo")
    r.sheet_view.showGridLines = False
    r.column_dimensions["A"].width = 3
    for letra, w in zip("BCDEFGHIJK", [10, 12, 14, 14, 9, 13, 11, 13, 13, 9]):
        r.column_dimensions[letra].width = w
    r["B2"] = "Permeabilidade por etapas de vazão"
    r["B2"].font = Font(size=14, bold=True, color=AZUL)
    info = [("Amostra", meta.get("Descrição da amostra")), ("Data do ensaio", meta.get("Data")), ("Operador", meta.get("Usuário")),
            ("Comprimento (cm)", p.comprimento_cm or meta.get("Comprimento (cm)")), ("Diâmetro (cm)", p.diametro_cm or meta.get("Diâmetro (cm)")),
            ("Viscosidade (cP)", p.viscosidade_cp or meta.get("Viscosidade (cP)")), ("Fluido", meta.get("Fluido de tratamento")),
            ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M"))]
    for i, (k, v) in enumerate(info, start=4):
        r.cell(row=i, column=2, value=k).font = Font(bold=True, color="475569")
        r.cell(row=i, column=4, value=v).alignment = Alignment(horizontal="left")
    cab_r = ["Etapa", "Q (cm³/min)", "Início estável (s)", "Fim estável (s)", "Pontos", "ΔP médio (psi)", "Desvio (psi)",
             "Volume inj. (cm³)", "k (mD)", "Usada"]
    lin0 = 4 + len(info) + 1
    for c, t in enumerate(cab_r, start=2):
        cel = r.cell(row=lin0, column=c, value=t)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r.row_dimensions[lin0].height = 30
    for i, e in enumerate(sorted(p.etapas, key=lambda x: x.ordem), start=lin0 + 1):
        vals = [e.ordem, e.q, e.j_inicio, e.j_fim, e.n_pontos, e.dp_medio, e.dp_desvio, e.volume_injetado_cm3, e.k_md, "sim" if e.usada else "não"]
        fmts = ["0", "0.00", "0.0", "0.0", "0", "0.000", "0.000", "0.00", "0.00", "@"]
        for c, (v, f) in enumerate(zip(vals, fmts), start=2):
            cel = r.cell(row=i, column=c, value=v)
            cel.number_format = f
            cel.alignment = Alignment(horizontal="center")
            cel.border = Border(bottom=FINO)
    # k final
    usadas = [e for e in p.etapas if e.usada and e.dp_medio]
    lin = lin0 + len(p.etapas) + 2
    L, D, mu = p.comprimento_cm, p.diametro_cm, p.viscosidade_cp
    if len(usadas) >= 2 and L and D and mu:
        a, b, r2, a_l, _, a_0 = ajuste_darcy([(e.q, e.dp_medio) for e in usadas], p.ajuste)
        area = math.pi * D * D / 4
        k_de = lambda incl: CONSTANTE_DARCY * mu * L / (area * incl)  # noqa: E731
        k_et = [e.k_md for e in usadas if e.k_md]
        nome_aj = "reta com intercepto" if p.ajuste != "origem" else "reta pela origem"
        res = [(f"k pelo ajuste ΔP × Q — {nome_aj} (mD)", k_de(a), "0.00"),
               (f"Reta: ΔP = {a:.4f}·Q {'−' if b < 0 else '+'} {abs(b):.4f}   ·   R²", r2, "0.0000"),
               ("k médio das etapas (mD)", sum(k_et) / len(k_et) if k_et else None, "0.00"),
               ("k pela outra reta (comparação, mD)", k_de(a_0 if p.ajuste != "origem" else a_l), "0.00"),
               ("k usado na amostra", p.metodo_k, "@")]
        for i, (k, v, f) in enumerate(res):
            r.cell(row=lin + i, column=2, value=k).font = Font(bold=True, color="475569")
            cel = r.cell(row=lin + i, column=8, value=v)
            cel.number_format = f
            if i == 0:
                cel.font = Font(bold=True, size=12, color="C8102E")
    r.cell(row=lin + 6, column=2, value="k = 244,93 · μ · L / (A · inclinação de ΔP × Q)   [mD, cP, cm, cm², psi/(cm³/min)] — igual a Q·μ·L/(A·ΔP) em SI ÷ 9,86923e-16").font = Font(italic=True, color="64748B")

    _dados_por_vazao(r, lin + 9, p, meta, ks, dados, i_t)

    _impressao(r, paisagem=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CORES_VAZAO = ["1F4E9C", "C8102E", "1B7F4B", "8E3FA8", "B5651D", "0F7C8C", "5B5B5B"]


def _coluna_volume(meta: dict, ks: list, dados: list) -> tuple[int, str]:
    """Volume bombeado do fluido de tratamento: a garrafa com o mesmo nome do 'Fluido de tratamento';
    se não achar, a garrafa cujo volume bombeado mais varia. Devolve (-1, ...) se não houver nenhuma."""
    fluido = chave(meta.get("Fluido de tratamento"))
    candidatas = []
    for g in range(1, 9):
        jn, jv = f"nome g{g}", f"volume bomb parcial g{g} cm"
        if jv not in ks:
            continue
        iv = ks.index(jv)
        vals = [l[iv] for l in dados if isinstance(l[iv], (int, float))]
        variacao = (max(vals) - min(vals)) if vals else 0
        nome = next((l[ks.index(jn)] for l in dados if jn in ks and l[ks.index(jn)]), None)
        if fluido and nome and chave(nome) == fluido and variacao > 0:
            return iv, f"garrafa {g} ({nome})"
        candidatas.append((variacao, iv, f"garrafa {g}" + (f" ({nome})" if nome else "")))
    candidatas.sort(reverse=True)
    if candidatas and candidatas[0][0] > 0:
        return candidatas[0][1], candidatas[0][2]
    return -1, "Q × tempo"


def _dados_por_vazao(r, lin: int, p: PedidoPlanilha, meta: dict, ks: list, dados: list, i_t: int):
    """Bloco com TODAS as leituras extraídas (trecho estável) de cada vazão, lado a lado:
    para cada vazão, 3 colunas — Tempo (s), ΔP (psi) e Volume (cm³)."""
    i_sel = next((j for j, k in enumerate(ks) if k.startswith("selecao do holder")), -1)
    h2 = i_sel >= 0 and sum("2" in str(l[i_sel] or "") for l in dados) > sum("1" in str(l[i_sel] or "") for l in dados)
    i_dp = ks.index("dph2 psi" if h2 else "dph1 psi")
    i_q = ks.index("q cm min")
    i_v, origem_v = _coluna_volume(meta, ks, dados)

    r.cell(row=lin, column=2, value="Dados extraídos de cada vazão (trecho estável)").font = Font(size=12, bold=True, color=AZUL)
    r.cell(row=lin + 1, column=2,
           value=f"ΔP do holder {'2' if h2 else '1'} · Volume = volume bombeado acumulado da {origem_v}").font = Font(italic=True, color="64748B")
    topo = lin + 3
    etapas = sorted(p.etapas, key=lambda x: x.ordem)
    for n, e in enumerate(etapas):
        c0 = 2 + n * 4  # 3 colunas de dados + 1 de folga
        cor = CORES_VAZAO[n % len(CORES_VAZAO)]
        r.merge_cells(start_row=topo, start_column=c0, end_row=topo, end_column=c0 + 2)
        tit = r.cell(row=topo, column=c0, value=f"Q = {e.q:g} cm³/min ({e.rotulo or f'Permeabilidade {e.ordem}'})")
        tit.font = Font(bold=True, color="FFFFFF")
        tit.fill = PatternFill("solid", fgColor=cor)
        tit.alignment = Alignment(horizontal="center")
        for k, t in enumerate(["Tempo (s)", "ΔP (psi)", "Volume (cm³)"]):
            cel = r.cell(row=topo + 1, column=c0 + k, value=t)
            cel.font = Font(bold=True, color=cor)
            cel.alignment = Alignment(horizontal="center")
            cel.border = Border(bottom=Side(style="medium", color=cor))
            letra = get_column_letter(c0 + k)
            r.column_dimensions[letra].width = max(r.column_dimensions[letra].width or 0, 12)
        linhas = [l for l in dados if e.j_inicio <= l[i_t] <= e.j_fim]
        acum, t_ant = 0.0, None
        for m, l in enumerate(linhas):
            if i_v >= 0:
                vol = l[i_v]
            else:  # sem garrafa: integra Q (cm³/min) no tempo (s)
                if t_ant is not None and isinstance(l[i_q], (int, float)):
                    acum += l[i_q] * (l[i_t] - t_ant) / 60
                vol, t_ant = acum, l[i_t]
            for k, (v, f) in enumerate([(l[i_t], "0.0"), (l[i_dp], "0.000"), (vol, "0.00")]):
                cel = r.cell(row=topo + 2 + m, column=c0 + k, value=v)
                cel.number_format = f
                cel.alignment = Alignment(horizontal="center")
    r.print_title_rows = f"{topo}:{topo + 1}"


def _impressao(ws, paisagem: bool = True):
    """Imprime numa página de largura (A4), repetindo o cabeçalho da tabela."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if paisagem:
        ws.print_title_rows = "2:2"
