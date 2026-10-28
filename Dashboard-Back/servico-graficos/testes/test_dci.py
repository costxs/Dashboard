"""Equipamento DCI no serviço de gráficos: leitor, planilha de etapas e acidificação (arquivos fictícios)."""
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

AQUI = Path(__file__).parent


@pytest.fixture
def cliente(monkeypatch):
    monkeypatch.setenv("MODO_DEV", "1")
    from app import app
    return TestClient(app)


def test_dci_leitor_normaliza():
    from graficos.leitor import ler_arquivo
    tb = ler_arquivo((AQUI / "dci_permeabilidade_ficticio.xlsx").read_bytes())
    assert tb.equipamento == "DCI" and tb.meta["amostra"] == "TESTE-D"
    assert abs(tb.meta["diametro_cm"] - 2.54) < 1e-9 and abs(tb.meta["comprimento_cm"] - 5.08) < 1e-9
    assert tb.meta["data"] == "05/10/2026 09:00:00" and tb.meta["porosidade_pct"] == 15
    assert abs(tb.t(tb.linhas[100]) - 100) < 0.01                      # horas → s
    assert tb.cab[tb.idx["dp"]].startswith("Core B Differential Pressure Total")


def test_dci_vira_o_formato_do_jbv():
    from graficos.leitor import dci_como_jbv, eh_dci
    c = (AQUI / "dci_permeabilidade_ficticio.xlsx").read_bytes()
    assert eh_dci(c) and not eh_dci((AQUI / "permeabilidade_ficticio.xlsx").read_bytes())
    x = dci_como_jbv(c)
    assert x.cab[:3] == ["Tempo (s)", "dPH1 (psi)", "Q (cm³/min)"] and x.meta["Equipamento"] == "DCI"
    # a opção DP1 + DP2 troca o ΔP (o total inclui o DP3)
    soma = dci_como_jbv(c, "dp1_dp2")
    assert any(a[1] is not None and b[1] is not None and a[1] != b[1] for a, b in zip(x.linhas, soma.linhas))


def test_dci_planilha_etapas(cliente):
    from graficos.leitor import ler_arquivo
    tb = ler_arquivo((AQUI / "dci_permeabilidade_ficticio.xlsx").read_bytes())
    assert tb.equipamento == "DCI"
    etapas = [{"ordem": 1, "q": 1, "j_inicio": 300, "j_fim": 599.5}, {"ordem": 2, "q": 2, "j_inicio": 800, "j_fim": 999.5},
              {"ordem": 3, "q": 4, "j_inicio": 1150, "j_fim": 1299.5}]
    corpo = {"etapas": etapas, "comprimento_cm": 5.08, "diametro_cm": 2.54, "viscosidade_cp": 1}
    with open(AQUI / "dci_permeabilidade_ficticio.xlsx", "rb") as f:
        r = cliente.post("/permeabilidade/planilha-etapas", files={"arquivo": ("dci.xlsx", f)}, data={"pedido": json.dumps(corpo)})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    cab = [c.value for c in wb["Planilha1"][2]][1:]
    assert cab[0] == "Timestamp (hours)" and "Core B Differential Pressure Total (DPBtot) (psi)" in cab and cab[-1] == "Etapa"
    assert wb["Planilha1"].max_row - 2 == 300 + 200 + 150
    res = wb["Resumo"]
    vals = {res.cell(row=i, column=2).value: res.cell(row=i, column=4).value for i in range(4, 14)}
    assert vals["Equipamento"] == "DCI" and vals["Amostra"] == "TESTE-D"


def test_dci_acidificacao_planilha_e_graficos(cliente):
    pedido = {"comprimento_cm": 5.08, "diametro_cm": 2.54, "porosidade_pct": 15, "volume_morto_cm3": 0.4, "temperatura_rotulo": "T50"}
    with open(AQUI / "dci_acidificacao_ficticio.xlsx", "rb") as f:
        r = cliente.post("/acidificacao/planilha", files={"arquivo": ("dci.xlsx", f)}, data={"pedido": json.dumps(pedido)})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ini = {wb["0- Initial Data"].cell(row=i, column=1).value: wb["0- Initial Data"].cell(row=i, column=2).value for i in range(1, 20)}
    assert ini["Equipamento"] == "DCI" and ini["Sample"] == "TESTE-D" and ini["PVbt (fim do ácido)"].startswith("='2 - Acidizing'!E")
    assert wb["2 - Acidizing"]["F3"].value == "Acumulador 2 @0.5 ml/min T50"
    for qual in ("geral", "pv", "apos"):
        with open(AQUI / "dci_acidificacao_ficticio.xlsx", "rb") as f:
            g = cliente.post("/acidificacao/grafico", files={"arquivo": ("dci.xlsx", f)}, data={"pedido": json.dumps(pedido), "qual": qual})
        assert g.status_code == 200 and g.content[:4] == b"\x89PNG"


def test_dci_sem_breakthrough():
    from graficos.acidificacao import Dados, avaliar_breakthrough, detectar_fases, gerar_planilha, PedidoAcid
    c = (AQUI / "dci_acidificacao_sem_bt_ficticio.xlsx").read_bytes()
    d = Dados(c)
    fs = detectar_fases(d)
    for f in fs:
        f.papel = "pre_fluxo" if f.ordem == 1 else "acido" if f.ordem == 2 else "pos_fluxo"
    houve, razao = avaliar_breakthrough(d, fs)
    assert houve is False and razao > 4
    p = PedidoAcid(comprimento_cm=5.08, diametro_cm=2.54, porosidade_pct=15, volume_morto_cm3=0.4, fases=fs)
    wb = load_workbook(io.BytesIO(gerar_planilha(c, p)))
    ini = {wb["0- Initial Data"].cell(row=i, column=1).value: wb["0- Initial Data"].cell(row=i, column=2).value for i in range(1, 20)}
    assert ini["PVbt (fim do ácido)"] == "sem breakthrough"


def test_tempo_em_horas_no_grafico_de_permeabilidade(cliente):
    p = json.loads((AQUI / "exemplo_pedido.json").read_text())
    p["config"].update(formato="svg", modo="combinado", tempo_em_horas=True)
    r = cliente.post("/permeabilidade/grafico", json=p)
    assert r.status_code == 200, r.text
    assert "Tempo (h)" in r.text and "Tempo (s)" not in r.text


def test_dci_tempo_em_horas_na_planilha_de_acidificacao(cliente):
    from graficos.acidificacao import Dados, detectar_fases
    c = (AQUI / "dci_acidificacao_ficticio.xlsx").read_bytes()
    d = Dados(c)
    assert d.div_t == 3600 and d.rot_t == "h"
    acido = max(detectar_fases(d), key=lambda f: 0)  # só para ter a janela de uma fase
    pedido = {"comprimento_cm": 5.08, "diametro_cm": 2.54, "porosidade_pct": 15, "volume_morto_cm3": 0.4}
    r = cliente.post("/acidificacao/planilha", files={"arquivo": ("dci.xlsx", c)}, data={"pedido": json.dumps(pedido)})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["2 - Acidizing"]
    # coluna do tempo na unidade do equipamento e a fórmula do volume converte horas em minutos
    assert ws["B5"].value < 5, "o tempo deveria estar em horas (o ensaio fictício dura < 1 h), não em segundos"
    assert str(ws["C5"].value).endswith("*60")
    # o PV não muda com a unidade: o mesmo valor que o cálculo em segundos
    t0_s = ws["B5"].value * 3600
    assert any(abs(f.t_inicio - t0_s) < 5 or f.t_inicio <= t0_s <= f.t_fim for f in detectar_fases(d))


def test_jbv_continua_em_segundos(cliente):
    from graficos.acidificacao import Dados
    d = Dados((AQUI / "acidificacao_ficticio.xlsx").read_bytes())
    assert d.div_t == 1 and d.rot_t == "s"
