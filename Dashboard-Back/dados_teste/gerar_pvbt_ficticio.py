"""Gera pvbt_ficticio.xlsx com a MESMA ESTRUTURA da planilha "Banco de dados de PVbt —
LCPETRO/UFPA" (abas LEIA-ME, dicionario, experimento, amostras, sistemas_fluidos,
fluidos_prop, referencias, qualidade, listas; uma Tabela do Excel por aba), mas com
dados INVENTADOS e erros de propósito para testar o importador."""
from datetime import datetime
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

EXP_CAB = ['id_experimento', 'fonte', 'id_referencia', 'id_amostra', 'litologia', 'd_nominal_in', 'l_nominal_in',
  'dimensao_nominal_in', 'diametro_cm', 'comprimento_cm', 'porosidade_pct', 'faixa_porosidade', 'volume_aparente_cm3',
  'volume_poroso_cm3', 'k_inicial_mD', 'faixa_permeabilidade', 'acido', 'conc_valor', 'conc_unidade', 'conc_pct', 'conc_M',
  'aditivo', 'aditivo_valor', 'aditivo_unidade', 'tratamento', 'temperatura_C', 'temperatura_nominal_C', 'vazao_ml_min',
  'vazao_nominal_ml_min', 'vel_darcy_cm_min', 'vel_intersticial_cm_min', 'vel_intersticial_nominal_cm_min', 'pvbt',
  'pvbt_condicao', 'pvbt_classe', 'comprimento_wormhole_cm', 'comprimento_wormhole_uct_cm', 'lwh_referencia',
  'volume_poroso_wh_cm3', 'pvbt_por_darcy', 'pvbt_por_uct', 'fluido_saturacao', 'tipo_fluido_saturacao', 'metodo_saturacao',
  'forma_saturacao', 'obtencao_dado', 'p_confinamento_psi', 'p_backpressure_psi', 'vel_intersticial_fonte', 'k_fluxo_mD',
  'breakthrough', 'temperatura_obs', 'tipo_experimento', 'status', 'data', 'ano', 'projeto', 'estudo', 'equipamento',
  'id_sistema', 'rotulo_original', 'observacoes']

def exp(id_, fonte, ref, am, acido, cv, cu, temp, vazao, pvbt, cond, **k):
    d = dict(id_experimento=id_, fonte=fonte, id_referencia=ref, id_amostra=am, acido=acido, conc_valor=cv, conc_unidade=cu,
             temperatura_C=temp, vazao_ml_min=vazao, pvbt=pvbt, pvbt_condicao=cond, pvbt_classe='normal',
             tratamento='Ácido inorgânico', tipo_experimento='Monocore', status='Finalizada')
    d.update(k)
    # colunas calculadas: a planilha real traz valores; aqui um valor qualquer (o importador ignora)
    d.update(litologia='(calc)', volume_poroso_cm3=0, conc_M=0)
    return [d.get(c) for c in EXP_CAB]

EXP = [
  exp('LCP-9001', 'LCPETRO', 'REF-LCPX', 'IL901', 'HCl', 15, 'wt%', 45, 0.18, 6.45, 'completo', id_sistema='SF92',
      aditivo='IC', aditivo_valor=1, aditivo_unidade='vol%', fluido_saturacao='AUP', tipo_fluido_saturacao='aquoso',
      metodo_saturacao='vácuo/saturador', obtencao_dado='bancada', k_fluxo_mD=1.31, data=datetime(2023, 11, 12), ano=2023,
      projeto='Projeto X', estudo='PVBt - HCl 15 wt% - 45 °C', equipamento='Core A', rotulo_original='HCl [15 wt%] @0.18 T45',
      vel_intersticial_fonte=0.1012, observacoes='# Não ocorreu Bt | camisa furou'),
  exp('LCP-9002', 'LCPETRO', 'REF-LCPX', 'IL902', 'HCl', 15, 'wt%', 45, 0.5, 2.8, 'parcial', id_sistema='SF92',
      comprimento_wormhole_cm=3.9, comprimento_wormhole_uct_cm=4.29, lwh_referencia='microCT', breakthrough='nao',
      obtencao_dado='bancada', data=datetime(2023, 11, 20), ano=2023),
  exp('LCP-9003', 'LCPETRO', 'REF-LCPX', 'IL900', 'HCl', 15, 'wt%', 45, 1.0, 1.9, 'completo', id_sistema='SF92',
      breakthrough='sim', obtencao_dado='bancada', data=datetime(2023, 12, 1), ano=2023),
] + [
  exp(f'LIT-900{i}', 'Literatura', 'REF-LX1', f'LIT-A0{i}', 'HCl', 15, 'wt%', 'Ambiente' if i == 1 else None, v, p, 'completo',
      id_sistema='SF91', obtencao_dado='grafico', temperatura_obs=None if i == 1 else 'Ambiente', ano=2001)
  for i, (v, p) in enumerate([(0.5, 3.1), (1, 1.2), (2, 0.8), (5, 0.9), (10, 1.4)], start=1)
] + [
  exp('LIT-9006', 'Literatura', 'REF-LX1', 'LIT-A99', 'HCl', 15, 'wt%', 25, 2, 0.7, 'completo', id_sistema='SF91'),  # amostra inexistente
  exp('LIT-9007', 'Literatura', 'REF-LX1', 'LIT-A06', 'HCl', 0.5, 'M', 25, 2, 4.0, '*', id_sistema='SF93'),          # condição inválida
  exp('LIT-9003', 'Literatura', 'REF-LX1', 'LIT-A03', 'HCl', 15, 'wt%', 25, 2, 0.8, 'completo'),                     # id repetido
  exp('PB-9001', 'Petrobras', 'REF-PBX', 'PB-AA-01', 'HCl', 15, 'wt%', 93.3, 1, 1.6, 'completo', id_sistema='SF91',
      obtencao_dado='tabela', p_confinamento_psi=3000, p_backpressure_psi=1500, ano=2019),
  exp('PB-9002', 'Petrobras', 'REF-PBX', 'PB-AA-02', 'HCl', 15, 'wt%', 93.33, 3, 0.95, 'completo', id_sistema='SF91',
      obtencao_dado='tabela', p_confinamento_psi=3000, p_backpressure_psi=1500, ano=2019),
]

AMO_CAB = ['id_amostra', 'fonte', 'litologia', 'd_nominal_in', 'l_nominal_in', 'diametro_cm', 'comprimento_cm', 'porosidade_pct',
           'volume_poroso_medido_cm3', 'k_inicial_mD', 'k_metodo', 'ca_pct', 'mg_pct']
AMO = [
  ['IL901', 'LCPETRO', 'Indiana Limestone', 1.5, 3, 3.79, 7.65, 15.77, 13.6, 2.25, 'gás', None, None],
  ['IL902', 'LCPETRO', 'Indiana', 1.5, 3, 3.8, 7.6, 16.1, 20.0, 3.0, 'gás', None, None],   # VP medido diverge
  ['IL900', 'LCPETRO', 'Indiana Limestone', 1.5, 3, 3.8, 7.6, 17.0, None, 3.1, 'gás', None, None],
] + [[f'LIT-A0{i}', 'Literatura', 'Indiana limestone', 1, 6, 2.54, 15.24, 15.0, None, 5.0, 'não informado', 39.2, 0.3] for i in range(1, 7)] + [
  ['PB-AA-01', 'Petrobras', 'Coquina', 1.5, 3, 3.81, 7.62, 18.5, None, 40.0, 'gás', None, None],
  ['PB-AA-02', 'Petrobras', 'Coquina', 1.5, 3, 3.81, 7.62, 19.2, None, 55.0, 'gás', None, None],
]
SIS_CAB = ['id_sistema', 'nome_sistema', 'componente_1', 'valor_1', 'unidade_1', 'componente_2', 'valor_2', 'unidade_2',
           'componente_3', 'valor_3', 'unidade_3', 'componente_4', 'valor_4', 'unidade_4']
SIS = [['SF91', 'HCl 15 wt%', 'HCl', 15, 'wt%'] + [None] * 9,
       ['SF92', 'HCl 15 wt% + IC', 'HCl', 15, 'wt%', 'IC', 1, 'vol%'] + [None] * 6,
       ['SF93', 'HCl 0.5M', 'HCl', 0.5, 'M'] + [None] * 9]
FLU = [['HCl', 36.46, 0.00494, 'Ácido clorídrico (dado fictício para teste)']]
REF = [['REF-LCPX', 'interna', 'LCPETRO/UFPA', 2023, 'Ensaios de coreflooding (teste)', 'LCPetro'],
       ['REF-LX1', 'artigo', 'Autor Inventado at al', 2001, 'Artigo fictício de wormholing', 'Inventado 2001'],
       ['REF-PBX', 'relatorio_petrobras', 'Petrobras (teste)', 2019, 'Relatório fictício', 'PB 2019']]
QUA = [['LIT-9005', 'LIT-A05', 'pvbt', '1.4', None, 'valor atípico, conferir', None, None]]
LIS = {'fonte': ['LCPETRO', 'Literatura', 'Petrobras'], 'tratamento': ['Ácido inorgânico', 'Espuma ácida'],
       'conc_unidade': ['wt%', 'M'], 'aditivo_unidade': ['vol%', 'wt%', 'ml'],
       'pvbt_condicao': ['completo', 'parcial', 'nao_atingido', 'nao_calculado'], 'pvbt_classe': ['normal', 'otimo'],
       'obtencao_dado': ['bancada', 'tabela', 'grafico', 'texto'], 'breakthrough': ['sim', 'nao'], 'lwh_referencia': ['microCT', 'darcy'],
       'status': ['Finalizada', 'Agendado', 'Em andamento']}

def aba(wb, nome, cab, linhas, tabela):
    ws = wb.create_sheet(nome)
    ws.append(cab)
    for l in linhas: ws.append(l)
    t = Table(displayName=tabela, ref=f"A1:{ws.cell(row=1, column=len(cab)).column_letter}{1 + max(1, len(linhas))}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

wb = Workbook()
ws = wb.active; ws.title = 'LEIA-ME'
for l in ['Banco de dados de PVbt — ARQUIVO FICTÍCIO PARA TESTE', '', 'Mesma estrutura da planilha real; dados inventados.']: ws.append([l])
aba(wb, 'dicionario', ['tabela', 'campo', 'tipo', 'unidade', 'valores_aceitos', 'descricao'],
    [['experimentos', 'id_experimento', 'texto', None, 'LCP-####; LIT-####; PB-####', 'Identificador único do ensaio']], 'Tabela2')
aba(wb, 'experimento', EXP_CAB, EXP, 'Tabela1')
aba(wb, 'amostras', AMO_CAB, AMO, 'Tabela3')
aba(wb, 'sistemas_fluidos', SIS_CAB, SIS, 'Tabela4')
aba(wb, 'fluidos_prop', ['componente', 'massa_molar_g_mol', 'coef_densidade', 'observacao'], FLU, 'Tabela5')
aba(wb, 'referencias', ['id_referencia', 'tipo', 'autor', 'ano', 'titulo', 'rotulo_original'], REF, 'Tabela6')
aba(wb, 'qualidade', ['id_experimento', 'id_amostra', 'campo', 'valor_registrado', 'valor_calculado', 'observacao', 'Coluna1', 'Coluna2'], QUA, 'Tabela7')
n = max(len(v) for v in LIS.values())
aba(wb, 'listas', list(LIS), [[LIS[k][i] if i < len(LIS[k]) else None for k in LIS] for i in range(n)], 'Tabela8')
wb.save('pvbt_ficticio.xlsx')
print('ok: pvbt_ficticio.xlsx')
