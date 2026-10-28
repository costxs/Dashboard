"""Acidificação (equipamento JBV): a mesma pasta de trabalho e os mesmos gráficos que o laboratório monta no Origin.

Planilha gerada (mesmos nomes de aba do laboratório; colunas no formato do Origin, com as linhas
"long name / units / comments / fx" e os dados a partir da linha 5):
  0- Initial Data                  amostra, sistema ácido, vazão, D, L, φ, volume morto e as FÓRMULAS de
                                   volume poroso, área e velocidade intersticial
  1 - Water Injection - Before ac  Time × DP de cada fase antes do ácido (ex.: KCl 2% 2.0 ml/min)
  2 - Acidizing                    Time · Volume injetado Total · Injected Volume · Pore Volume (PV) · DP (com fórmulas)
  3 - Experiment general View      Time · Flow rate · Back pressure · Confining pressure
  4 - Temperature                  Time · Temperature (estufa)
  5 - Water Injection - After aci  Time × DP de cada fase depois do ácido (ex.: UW 2.0 ml/min (after))
  Gráficos                         os 3 gráficos abaixo

Gráficos (padrão Origin):
  geral  Graph 0 — ΔP × tempo, cada fase numa cor, e a vazão tracejada no eixo da direita
  pv     Graph 3 — ΔP × PV do ácido; PV = (Q·(t − t₀)/60 − volume morto)/VP, linha tracejada em PV = 0 ("Dead Volume")
  apos   Graph 4 — ΔP × tempo das fases depois do ácido (escala ajustada)

As fases normalmente vêm do gerencial (o mesmo tratamento que a pessoa conferiu na tela). Se não vierem,
são detectadas aqui pela garrafa cujo "Volume atual" está caindo + a vazão.
"""
from __future__ import annotations

import io
import math
from typing import Literal, Optional

import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, field_validator

from .estilo import acabamento, estilo_origin
from .planilha_etapas import AZUL, chave, ler_tabela_leituras

CORES_PRE = ["#0000FF", "#000080", "#008080"]
COR_ACIDO = "#FF0000"
CORES_POS = ["#000000", "#FF00FF", "#808000", "#800080", "#008000"]
COR_VAZAO = "#FF8000"


class FaseAcid(BaseModel):
    ordem: int
    papel: Literal["pre_fluxo", "acido", "pos_fluxo", "outro"]
    q: float
    t_inicio: float
    t_fim: float
    fluido: Optional[str] = None
    legenda: Optional[str] = None        # padrão: "KCl 2% 2.0 ml/min", "BLEND @2 ml/min", "AUP 2.0 ml/min (after)"
    cor: Optional[str] = None


class ConfigAcid(BaseModel):
    formato: Literal["png", "svg", "pdf"] = "png"
    dpi: int = Field(300, ge=72, le=1200)
    fonte: str = "Arial"
    tamanho_fonte: float = Field(12, ge=6, le=40)
    espessura_linha: float = Field(2.0, gt=0, le=8)
    mostrar_pvbt: bool = True


class PedidoAcid(BaseModel):
    amostra: Optional[str] = None
    sistema_acido: Optional[str] = None
    comprimento_cm: float = Field(gt=0)
    diametro_cm: float = Field(gt=0)
    porosidade_pct: float = Field(gt=0, lt=100)
    volume_morto_cm3: float = Field(0, ge=0)
    temperatura_rotulo: Optional[str] = None   # ex.: "T40" (vai na legenda do ácido, como no laboratório)
    aba: Optional[str] = None                   # aba que o gerencial leu (a mesma da tela)
    volume_poroso_cm3: Optional[float] = Field(None, gt=0)   # VP já calculado (gráficos do que foi gravado); senão, das dimensões
    houve_breakthrough: Optional[bool] = None   # o gerencial manda o que a pessoa confirmou; vazio = o serviço confere
    dp_dci: Literal["total", "dp1_dp2"] = "total"   # só DCI: ΔP = DP1+DP2+DP3 (total) ou só DP1+DP2
    fases: list[FaseAcid] = []
    config: ConfigAcid = ConfigAcid()

    @field_validator("volume_morto_cm3", mode="before")
    @classmethod
    def _morto_vazio_e_zero(cls, v):
        # campo em branco no gerencial chega como null: sem volume morto = 0
        return 0 if v is None or v == "" else v


# --------------------------------------------------------------------------- leitura
class Dados:
    # o arquivo do DCI é convertido para o formato do JBV (graficos/leitor.py); estes dois campos dizem de onde veio
    equipamento = "JBV"
    nome_garrafa = "garrafa"       # no DCI, "acumulador"
    # o tempo do DCI nasce em horas: o serviço calcula em segundos, mas mostra (gráficos e planilhas) na unidade do equipamento
    div_t = 1.0                    # segundos por unidade mostrada (DCI: 3600)
    rot_t = "s"

    def __init__(self, conteudo: bytes, aba: str | None = None, dp_dci: str = "total"):
        self.aba, self.meta, cab, linhas = ler_tabela_leituras(conteudo, aba, dp_dci)
        if self.meta.get("Equipamento") == "DCI":
            self.equipamento, self.nome_garrafa = "DCI", "acumulador"
            self.div_t, self.rot_t = 3600.0, "h"
        self.ks = [chave(v) for v in cab]
        i_t = self.ks.index("tempo s")
        self.linhas = [l for l in linhas if l and isinstance(l[i_t], (int, float))]
        sel = self._col(lambda k: k.startswith("selecao do holder"))
        h2 = sel >= 0 and sum("2" in str(l[sel] or "") for l in self.linhas) > sum("1" in str(l[sel] or "") for l in self.linhas)
        h = "h2" if h2 else "h1"
        self.c = {
            "t": i_t, "q": self._col(lambda k: k.startswith("q cm")), "dp": self._col(lambda k: k == f"dp{h} psi"),
            "pb": self._col(lambda k: k == f"pb{h} psi"), "pc": self._col(lambda k: k == "pc psi"),
            "temp": self._col(lambda k: k.startswith("t estufa")),
        }
        if self.c["dp"] < 0 or self.c["q"] < 0:
            raise ValueError(f"colunas dP{h.upper()} (psi) e Q (cm³/min) não encontradas")
        self.garrafas = {}
        for j, k in enumerate(self.ks):
            if k.startswith("volume atual g"):
                n = int("".join(ch for ch in k.split("volume atual g")[1] if ch.isdigit())[:1])
                nome = self._col(lambda x, n=n: x == f"nome g{n}")
                self.garrafas[n] = (j, nome)

    def _col(self, f) -> int:
        return next((j for j, k in enumerate(self.ks) if f(k)), -1)

    def v(self, l, nome):
        j = self.c[nome]
        x = l[j] if j >= 0 else None
        return x if isinstance(x, (int, float)) else None

    def janela(self, f: FaseAcid):
        # folga de 1 ms: os tempos vêm do mesmo arquivo, mas passam por JSON no caminho
        return [l for l in self.linhas if f.t_inicio - 1e-3 <= l[self.c["t"]] <= f.t_fim + 1e-3]

    def exigir_pontos(self, f: FaseAcid):
        """Erro claro quando a fase pedida não tem leituras no arquivo lido aqui (aba ou tempos diferentes da tela)."""
        if self.janela(f):
            return
        ts = [l[self.c["t"]] for l in self.linhas]
        lido = f"{ts[0]:g} a {ts[-1]:g} s, {len(ts)} leituras" if ts else "nenhuma leitura com tempo numérico"
        raise ValueError(f"a fase de ácido ({f.t_inicio:g} a {f.t_fim:g} s) não tem leituras no arquivo lido pelo serviço "
                         f"(aba '{self.aba}': {lido}). Confira se a aba e o arquivo são os mesmos da tela")


def detectar_fases(d: Dados, lim: float = 0.005, tol: Optional[float] = None, min_pontos: Optional[int] = None) -> list[FaseAcid]:
    """Garrafa (JBV) ou acumulador (DCI) cujo 'Volume atual' caiu + vazão. Papel: o ácido é a fase de maior queda de ΔP.
    No DCI a vazão tem ruído (usa a mediana móvel) e, em vazão baixa, o volume muda pouco entre duas leituras
    (olha algumas leituras para trás); as fases precisam de mais leituras."""
    dci = d.equipamento == "DCI"
    tol = tol if tol is not None else (0.10 if dci else 0.02)
    min_pontos = min_pontos if min_pontos is not None else (30 if dci else 3)
    ls, blocos = d.linhas, []
    qraw = [d.v(l, "q") or 0 for l in ls]
    h = 7 if dci else 0
    qs = [sorted(qraw[max(0, i - h):i + h + 1])[len(qraw[max(0, i - h):i + h + 1]) // 2] for i in range(len(ls))]
    zero = 0.02 if dci else 0
    voltas = (1, 5, 20, 60) if dci else (1,)
    for i in range(1, len(ls)):
        q = qs[i]
        if q <= zero:
            continue
        g_at = None
        for volta in voltas:
            a_l, dv = ls[max(0, i - volta)], lim
            for g, (j, _) in d.garrafas.items():
                a, b = a_l[j], ls[i][j]
                if isinstance(a, (int, float)) and isinstance(b, (int, float)) and a - b > dv:
                    g_at, dv = g, a - b
            if g_at is not None:
                break
        if g_at is None and blocos and blocos[-1]["i1"] == i - 1:
            g_at = blocos[-1]["g"]
        if g_at is None:
            continue
        u = blocos[-1] if blocos else None
        if u and u["g"] == g_at and u["i1"] == i - 1 and abs(q - u["q"]) <= tol * u["q"]:
            u["i1"] = i
        else:
            blocos.append({"g": g_at, "q": q, "i0": i, "i1": i})
    blocos = [b for b in blocos if b["i1"] - b["i0"] + 1 >= min_pontos]
    for b in blocos:                            # vazão nominal = mediana da fase
        v = sorted(qs[b["i0"]:b["i1"] + 1])
        b["q"] = round(v[len(v) // 2], 3 if v[len(v) // 2] < 1 else 2)   # sem o ruído do medidor (0,500319 -> 0,5)
    fases = []
    for k, b in enumerate(blocos):
        j_nome = d.garrafas[b["g"]][1]
        nome = next((ls[i][j_nome] for i in range(b["i0"], b["i1"] + 1) if j_nome >= 0 and ls[i][j_nome]), None) or f"{d.nome_garrafa.capitalize()} {b['g']}"
        dps = [d.v(ls[i], "dp") for i in range(b["i0"], b["i1"] + 1)]
        dps = [x for x in dps if x is not None]
        n = max(1, min(5, len(dps) // 3))
        queda = (sorted(dps[:n])[n // 2] / max(sorted(dps[-n:])[n // 2], 1e-9)) if dps else 1
        fases.append((queda, FaseAcid(ordem=k + 1, papel="outro", q=b["q"], fluido=str(nome),
                                      t_inicio=ls[b["i0"]][d.c["t"]], t_fim=ls[b["i1"]][d.c["t"]])))
    acido = max(fases, key=lambda x: x[0]) if fases else None
    if acido and acido[0] > 1.5:
        for _, f in fases:
            f.papel = "pre_fluxo" if f.ordem < acido[1].ordem else ("acido" if f.ordem == acido[1].ordem else "pos_fluxo")
    return [f for _, f in fases]


def _q_txt(q: float) -> str:
    return f"{q:.1f}" if q >= 0.95 else f"{q:.3g}"     # 2.0 · 3.0 · 0.081


def _legenda(f: FaseAcid, temp: Optional[str]) -> str:
    if f.legenda:
        return f.legenda
    nome = f.fluido or "?"
    if f.papel == "acido":
        return f"{nome} @{f.q:g} ml/min" + (f" {temp}" if temp else "")
    return f"{nome} {_q_txt(f.q)} ml/min" + (" (after)" if f.papel == "pos_fluxo" else "")


def avaliar_breakthrough(d: "Dados", fases: list[FaseAcid]) -> tuple[Optional[bool], Optional[float]]:
    """Compara o ΔP com o fluido de antes e de depois do ácido, na MESMA vazão (a mesma regra do gerencial).
    depois/antes < 0,5 → abriu caminho (breakthrough); ≥ 0,8 → não abriu; entre os dois → incerto (None)."""
    fa = next((f for f in fases if f.papel == "acido"), None)
    if not fa:
        return None, None
    mesma = lambda f: abs(f.q - fa.q) <= 0.1 * fa.q  # noqa: E731
    antes = next((f for f in sorted(fases, key=lambda x: -x.ordem) if f.ordem < fa.ordem and mesma(f)), None)
    depois = next((f for f in sorted(fases, key=lambda x: x.ordem) if f.ordem > fa.ordem and mesma(f)), None)
    if not antes or not depois:
        return None, None
    va = [x for x in (d.v(l, "dp") for l in d.janela(antes)) if x is not None]
    vd = sorted(x for x in (d.v(l, "dp") for l in d.janela(depois)) if x is not None)
    if not va or not vd:
        return None, None
    fim = sorted(va[int(len(va) * 0.8):])
    a, b = fim[len(fim) // 2], vd[len(vd) // 2]
    if a <= 0:
        return None, None
    r = b / a
    return (True if r < 0.5 else False if r >= 0.8 else None), r


def _cores(fases: list[FaseAcid]) -> dict[int, str]:
    c, ip, io_ = {}, 0, 0
    for f in sorted(fases, key=lambda x: x.ordem):
        if f.cor:
            c[f.ordem] = f.cor
        elif f.papel == "acido":
            c[f.ordem] = COR_ACIDO
        elif f.papel == "pos_fluxo":
            c[f.ordem] = CORES_POS[io_ % len(CORES_POS)]; io_ += 1
        else:
            c[f.ordem] = CORES_PRE[ip % len(CORES_PRE)]; ip += 1
    return c


# --------------------------------------------------------------------------- cálculos
def geometria(p: PedidoAcid) -> dict:
    a = math.pi * p.diametro_cm ** 2 / 4
    phi = p.porosidade_pct / 100
    vp = p.volume_poroso_cm3 or a * p.comprimento_cm * phi
    return {"area": a, "phi": phi, "vp": vp}


def serie_acido(d: Dados, f: FaseAcid, p: PedidoAcid) -> list[tuple[float, float, float, float, float]]:
    """(tempo, volume total, volume injetado, PV, ΔP) — igual à aba '2 - Acidizing' do Origin."""
    g = geometria(p)
    jan = d.janela(f)
    if not jan:
        return []
    t0 = jan[0][d.c["t"]]
    morto = p.volume_morto_cm3 or 0.0
    vp = g["vp"] if g["vp"] > 0 else 1.0
    out = []
    for l in jan:
        t = l[d.c["t"]]
        vt = f.q * (t - t0) / 60
        vi = vt - morto
        out.append((t, vt, vi, vi / vp, d.v(l, "dp")))
    return out


# --------------------------------------------------------------------------- gráficos
def _fmt(v, casas=2):
    return f"{v:.{casas}f}".replace(".", ",")


def _eixos_virgula(ax, casas_x=2, casas_y=2):
    ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: _fmt(v, casas_x)))
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: _fmt(v, casas_y)))


def grafico(conteudo: bytes, p: PedidoAcid, qual: Literal["geral", "pv", "apos"]) -> tuple[bytes, str]:
    d = Dados(conteudo, p.aba, p.dp_dci)
    fases = sorted(p.fases or detectar_fases(d), key=lambda x: x.ordem)
    return _grafico(d, fases, p, qual, p.config.formato, p.config.dpi)


def _grafico(d: Dados, fases: list[FaseAcid], p: PedidoAcid, qual: str, formato: str, dpi: int) -> tuple[bytes, str]:
    cfg = p.config
    cores = _cores(fases)
    lw = cfg.espessura_linha
    with estilo_origin(cfg.fonte, cfg.tamanho_fonte, 1.5, True):
        if qual == "geral":
            fig, ax = plt.subplots(figsize=(30 / 2.54, 10 / 2.54))
            for f in fases:
                jan = d.janela(f)
                ax.plot([l[d.c["t"]] / d.div_t for l in jan], [d.v(l, "dp") for l in jan], color=cores[f.ordem], lw=lw,
                        label=_legenda(f, p.temperatura_rotulo))
            ax2 = ax.twinx()
            ts = [l[d.c["t"]] / d.div_t for l in d.linhas]
            ax2.plot(ts, [d.v(l, "q") for l in d.linhas], color=COR_VAZAO, ls="--", lw=lw * 0.75, label="Flow rate")
            ax2.set_ylabel("Flow rate (ml/min)", fontweight="bold")
            ax.set_xlabel(f"Time ({d.rot_t})", fontweight="bold"); ax.set_ylabel("DP (psi)", fontweight="bold")
            dps = [(d.v(l, "dp") or 0) for f in fases for l in d.janela(f)]
            dpmax = max(dps) if dps else 1.0
            ax.set_xlim(left=0); ax.set_ylim(0, (dpmax * 2.1) if dpmax > 0 else 1.0)
            qs = [(d.v(l, "q") or 0) for l in d.linhas]
            qmax = max(qs) if qs else 1.0
            ax2.set_ylim(0, (qmax * 1.15) if qmax > 0 else 1.0)
            acabamento(ax); ax.tick_params(right=False)
            ax2.tick_params(direction="in", which="both")
            _eixos_virgula(ax); _eixos_virgula(ax2)
            h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
            ax.legend(h1 + h2, l1 + l2, loc="upper center", frameon=True, edgecolor="black", fancybox=False)
        elif qual == "pv":
            fa = next((f for f in fases if f.papel == "acido"), None)
            if not fa:
                raise ValueError("nenhuma fase marcada como ácido")
            d.exigir_pontos(fa)
            s = serie_acido(d, fa, p)
            fig, ax = plt.subplots(figsize=(16 / 2.54, 13 / 2.54))
            ax.plot([x[3] for x in s], [x[4] for x in s], color=cores[fa.ordem], lw=lw, label=_legenda(fa, p.temperatura_rotulo))
            ax.axvline(0, color="black", ls="--", lw=1.2)
            x0, x1 = min(x[3] for x in s), max(x[3] for x in s)
            ax.set_xlim(x0, x1 + 0.04 * (x1 - x0) if x1 > x0 else x0 + 1)
            dp_validos = [x[4] for x in s if x[4] is not None]
            dp_max = max(dp_validos) if dp_validos else 1.0
            ax.set_ylim(0, (dp_max * 1.2) if dp_max > 0 else 1.0)
            if (p.volume_morto_cm3 or 0) > 0:
                largo = -x0 > 0.3 * (x1 - x0) if x1 > x0 else False
                ax.annotate("Dead Volume" if largo else "← Dead Volume", (x0 / 2 if largo else 0, ax.get_ylim()[1] * 0.38),
                            textcoords="offset points", xytext=(0, 0) if largo else (6, 0),
                            ha="center" if largo else "left", fontweight="bold", fontfamily="serif")
            if cfg.mostrar_pvbt and s[-1][4] is not None:
                pvbt = s[-1][3]
                houve = p.houve_breakthrough if p.houve_breakthrough is not None else avaliar_breakthrough(d, fases)[0]
                texto_pvbt = f"PVbt = {_fmt(pvbt, 3)}" if houve is not False else f"Sem breakthrough\nPV injetado = {_fmt(pvbt, 3)}"
                ax.plot([pvbt], [s[-1][4]], marker="o", color="black", ms=6, ls="none")
                ax.annotate(texto_pvbt, (pvbt, s[-1][4]), textcoords="axes fraction", xytext=(0.97, 0.62),
                            ha="right", fontweight="bold", arrowprops={"arrowstyle": "->", "lw": 1, "color": "black"})
            ax.set_xlabel("Pore Volume (PV)", fontweight="bold"); ax.set_ylabel("DP (psi)", fontweight="bold")
            acabamento(ax); _eixos_virgula(ax)
            ax.legend(loc="upper left", frameon=True, edgecolor="black", fancybox=False)
        else:  # apos
            pos = [f for f in fases if f.papel == "pos_fluxo"]
            if not pos:
                raise ValueError("nenhuma fase depois do ácido")
            fig, ax = plt.subplots(figsize=(15 / 2.54, 11 / 2.54))
            for f in pos:
                jan = d.janela(f)
                ax.plot([l[d.c["t"]] / d.div_t for l in jan], [d.v(l, "dp") for l in jan], color=cores[f.ordem], lw=lw * 0.6,
                        label=_legenda(f, p.temperatura_rotulo))
            ax.set_xlabel(f"Time ({d.rot_t})"); ax.set_ylabel("DP (psi)")
            acabamento(ax); _eixos_virgula(ax)
            ax.legend(loc="upper right", frameon=True, edgecolor="black", fancybox=False)
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format=formato, dpi=dpi if formato == "png" else None, bbox_inches="tight")
        plt.close(fig)
    return buf.getvalue(), {"png": "image/png", "svg": "image/svg+xml", "pdf": "application/pdf"}[formato]


# --------------------------------------------------------------------------- planilha
CAB_ORIGIN = ["long name", "units", "comments", "fx"]


def _bloco(ws, col: int, nome: str, unid: str, coment: str = "", fx: str = ""):
    for r, v in enumerate([nome, unid, coment, fx], start=1):
        c = ws.cell(row=r, column=col, value=v or None)
        c.font = Font(bold=r == 1, italic=r == 4, color="FFFFFF" if r == 1 else "1F2937")
        if r == 1:
            c.fill = PatternFill("solid", fgColor=AZUL)
        elif r <= 3:
            c.fill = PatternFill("solid", fgColor="FFF7D6")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(col)].width = 14


def _aba_origin(wb, titulo: str):
    ws = wb.create_sheet(titulo[:31])
    for r, t in enumerate(CAB_ORIGIN, start=1):
        ws.cell(row=r, column=1, value=t).font = Font(bold=True, color="475569")
    ws.column_dimensions["A"].width = 11
    ws.freeze_panes = "B5"
    return ws


def gerar_planilha(conteudo: bytes, p: PedidoAcid) -> bytes:
    d = Dados(conteudo, p.aba, p.dp_dci)
    fases = sorted(p.fases or detectar_fases(d), key=lambda x: x.ordem)
    acido = next((f for f in fases if f.papel == "acido"), None)
    if not acido:
        raise ValueError("nenhuma fase marcada como ácido")
    d.exigir_pontos(acido)
    g = geometria(p)
    wb = Workbook()

    # 0- Initial Data (com fórmulas, como no Origin)
    ws = wb.active; ws.title = "0- Initial Data"
    pos = [f for f in fases if f.papel == "pos_fluxo"]
    linhas = [("Sample", p.amostra or d.meta.get("Descrição da amostra")), ("Acid System", p.sistema_acido or acido.fluido),
              ("Acid Flow Rate (cc/min)", acido.q), ("Diameter (cm)", p.diametro_cm), ("Length (cm)", p.comprimento_cm),
              ("Porosity", g["phi"]), ("Dead volume (ml)", p.volume_morto_cm3),
              ("Pore Volume (cm³)", "=PI()*((B4^2)/4)*B5*B6"), ("Transversal Area (cm²)", "=(PI()*(B4^2))/4"),
              ("Interstitial Velocity (cm/min)", "=B3/(B9*B6)")]
    linhas += [(f"UW Flow rate {i}", f.q) for i, f in enumerate(pos, start=1)]
    houve, razao = (p.houve_breakthrough, None) if p.houve_breakthrough is not None else avaliar_breakthrough(d, fases)
    ref_pv = f"='2 - Acidizing'!E{4 + len(d.janela(acido))}"
    linhas += [("PV de ácido injetado", ref_pv),
               ("PVbt (fim do ácido)", ref_pv if houve is not False else "sem breakthrough"),
               ("Equipamento", d.equipamento)]
    if razao is not None:
        linhas.append(("ΔP depois ÷ ΔP antes (mesma vazão)", round(razao, 3)))
    for r, (k, v) in enumerate(linhas, start=1):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, color="475569")
        c = ws.cell(row=r, column=2, value=v)
        if r in (8, 9, 10) or k.startswith("PV"):
            c.font = Font(bold=True, color="C8102E"); c.number_format = "0.00000"
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 14

    # 1 - antes do ácido: Time × DP por fase
    ws = _aba_origin(wb, "1 - Water Injection - Before ac")
    _pares(ws, d, [f for f in fases if f.papel == "pre_fluxo"], p)

    # 2 - Acidizing (com fórmulas)
    ws = _aba_origin(wb, "2 - Acidizing")
    for col, (n, u, c, fx) in enumerate([("Time", d.rot_t, "", ""), ("Volume injetado Total", "ml", "", "Q·(t − t₀)·60" if d.div_t > 1 else "Q·(t − t₀)/60"),
                                         ("Injected Volume", "ml", "", "Total − Dead volume"), ("Pore Volume (PV)", "", "", "Injected ÷ Pore Volume"),
                                         ("DP", "psi", _legenda(acido, p.temperatura_rotulo), "")], start=2):
        _bloco(ws, col, n, u, c, fx)
    ini = "'0- Initial Data'"
    for i, (t, _vt, _vi, _pv, dp) in enumerate(serie_acido(d, acido, p)):
        r = 5 + i
        ws.cell(row=r, column=2, value=t / d.div_t)
        # tempo em horas (DCI): horas × 60 = minutos; em segundos (JBV): ÷ 60
        conv = "*60" if d.div_t > 1 else "/60"
        ws.cell(row=r, column=3, value=f"={ini}!$B$3*(B{r}-$B$5){conv}").number_format = "0.00000"
        ws.cell(row=r, column=4, value=f"=C{r}-{ini}!$B$7").number_format = "0.00000"
        ws.cell(row=r, column=5, value=f"=D{r}/{ini}!$B$8").number_format = "0.00000"
        ws.cell(row=r, column=6, value=dp)

    # 3 - visão geral
    ws = _aba_origin(wb, "3 - Experiment general View")
    for col, (n, u, c) in enumerate([("Time", d.rot_t, ""), ("Flow rate", "ml/min", ""), ("Pressure", "psi", "Back pressure"),
                                     ("Pressure", "psi", "Confining pressure")], start=2):
        _bloco(ws, col, n, u, c)
    for i, l in enumerate(d.linhas):
        for col, k in enumerate(["t", "q", "pb", "pc"], start=2):
            ws.cell(row=5 + i, column=col, value=d.v(l, k) if k != "t" else l[d.c["t"]] / d.div_t)

    # 4 - temperatura
    ws = _aba_origin(wb, "4 - Temperature")
    _bloco(ws, 2, "Time", d.rot_t); _bloco(ws, 3, "Temperature", "ºC", "Stove")
    for i, l in enumerate(d.linhas):
        ws.cell(row=5 + i, column=2, value=l[d.c["t"]] / d.div_t); ws.cell(row=5 + i, column=3, value=d.v(l, "temp"))

    # 5 - depois do ácido
    ws = _aba_origin(wb, "5 - Water Injection - After aci")
    _pares(ws, d, pos, p)

    # Gráficos
    ws = wb.create_sheet("Gráficos")
    lin = 1
    for titulo, qual in [("Graph 0 - General View", "geral"), ("Graph 3 - Acidizing vs PV", "pv"), ("Graph 4 - Injection after", "apos")]:
        if qual == "apos" and not pos:
            continue
        ws.cell(row=lin, column=1, value=titulo).font = Font(bold=True, size=12, color=AZUL)
        png, _ = _grafico(d, fases, p, qual, "png", 110)
        img = XlImage(io.BytesIO(png))
        ws.add_image(img, f"A{lin + 1}")
        lin += int(img.height / 20) + 4

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pares(ws, d: Dados, fases: list[FaseAcid], p: PedidoAcid):
    for k, f in enumerate(fases):
        c0 = 2 + 2 * k
        _bloco(ws, c0, "Time", d.rot_t); _bloco(ws, c0 + 1, "DP", "psi", _legenda(f, p.temperatura_rotulo))
        for i, l in enumerate(d.janela(f)):
            ws.cell(row=5 + i, column=c0, value=l[d.c["t"]] / d.div_t)
            ws.cell(row=5 + i, column=c0 + 1, value=d.v(l, "dp"))
