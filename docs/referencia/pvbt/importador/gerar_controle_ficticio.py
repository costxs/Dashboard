"""Gera controle_ficticio.xlsx: mesma ESTRUTURA da planilha de controle de rochas
(aba coreflooding + aba Planilha2 de RDA, Tabelas do Excel começando na linha 2,
colunas Coluna1/2/3, conferência com data e nomes no título), mas com dados INVENTADOS
e com erros de propósito para testar o importador."""
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

CF_CAB = ['Coluna1', 'Código', 'Formação', 'Conferência 11/08/26 - Ana e Bruno', 'Permeabilidade Fabri. (mD)',
          'Dimensões (DxL) (in)', 'Diâmetro (cm)', 'Comprimento (cm)', 'Massa (g)', 'Data da Petrófisica',
          'Densidade Bulk (g/cm3)', 'Densidade de grão calculada (g/cm3)', 'Permeabilidade à Gás (mD)',
          'Permeabilidade em Fluxo (mD)', 'Porosidade Exp (%)', 'Densidade do grão (g/cm3)', 'Densidade aparente (g/cm3',
          'Massa apos acidificação', 'Porosidade apos acidificação', 'Status', 'Obs Petrofísica // Utilização', 'Coluna2', 'Coluna3']
CF = [
  # cod, formação, conf, kfab, dim, d, L, massa, data, bulk, graoCalc, kgas, kliq, phi, grao, aparente, massaApos, phiApos, status, obs
  ['CF101', 'Indiana Limestone', 'ok', 46114, '1.5x3', 3.79, 7.61, 189.88, 45075, 2.21, 2.21, 2.39, 0.97, 16.56, None, None, 185.2, 18.9, 'Utilizada', 'Acidificada'],
  ['CF102', 'Pink desert', 'ok', '70-115', '1.5x6', 3.81, 15.2, 380.1, 45100, 2.35, None, 95.0, 60.2, 12.4, 2.71, 2.36, None, None, 'Disponível', None],
  ['cf 103', 'Morro do chaves', 'não localizada', None, '1.5x3', 3.8, 7.6, 170.0, None, None, None, None, None, None, None, None, None, None, 'Não localizada', None],
  ['CF104', 'Silurian Dolomite', 'ok', 300, '1x2', None, None, 60.3, 45200, 2.5, 2.84, 210.0, None, 135, None, None, None, None, 'Reservada', None],
  ['CF102', 'Indiana', 'ok', None, '1.5x3', 3.8, 7.6, 190.0, None, None, None, None, None, None, None, None, None, None, 'Disponível', 'linha duplicada'],
  ['CF106', 'Jandaira', None, None, '(vazio)', 3.79, 7.5, 175.5, 45210, None, None, 1.2, None, 22.1, None, None, None, None, 'Quebrada', None],
  ['IL900', 'Indiana', 'ok', None, '1.5x3', 3.8, 7.6, 191.0, 45300, None, None, 3.1, None, 17.0, None, None, None, None, 'Utilizada', 'reimportada'],
]
RDA_CAB = ['Código', 'Formação', 'Conferência', 'Permeabilidade (mD)', 'Diâmetro (cm)', 'Comprimento (cm)',
           'Massa (g)', 'Porosidade Exp (%)', 'Utilização', 'Status']
RDA = [
  ['RDA101', 'Indiana', 'ok', 4.96, 3.81, 4.7, None, 16.72, 'TR - HCl [0,5M] - Indiana - 25 °C @100', 'Utilizada'],
  ['RDA102', 'Sillurian', 'ok', 50.1, 3.81, 1.2, 30.5, 12.3, 'TR - HCl [15%] - Sillurian - 60 °C @500', 'Disponível'],
  ['RDA103', 'Coquina', None, '16-20', 3.81, 1.0, 25.0, None, None, 'Disponível'],
  ['TST-001', 'Indiana', 'ok', '2 - 4 md', 3.81, 1.0, 26.0, 15.0, None, 'Disponível'],
]

def aba(ws, cab, linhas, nome):
    ws.append([])                       # linha 1 vazia, como na planilha real
    ws.append(cab)
    for l in linhas:
        ws.append(([None] + l + [None, None]) if cab[0] == 'Coluna1' else l)
    ref = f"A2:{ws.cell(row=2, column=len(cab)).column_letter}{2 + len(linhas)}"
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

wb = Workbook()
aba(wb.active, CF_CAB, CF, 'Tabela2'); wb.active.title = 'coreflooding'
aba(wb.create_sheet('Planilha2'), RDA_CAB, RDA, 'Tabela3')
wb.save('controle_ficticio.xlsx')
print('ok: controle_ficticio.xlsx')
